import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from PIL import Image, ImageTk
import mysql.connector
import os
import sys
import shutil
import uuid
import subprocess
from tkinter import filedialog
from datetime import datetime, date, timedelta
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HOST = "localhost"
USER = "root"
PASSWORD = "sqldecami"
DATABASE = "restaurante_vivaldi"


LOGO_PNG = "vivaldi_sf.png"
LOGO_ICO = "vivaldi.ico"


CREMA = "#fbf8f1"
CREMA_CLARO = "#fbf7f1"
BEIGE = "#d9bc8c"
VERDE = "#223d2d"
VERDE_OSCURO = "#12261c"
VINO = "#6b2333"
BLANCO = "#ffffff"
GRIS = "#6f6f6f"
GRIS_CLARO = "#e8e3da"


def conectar():

    try:

        return mysql.connector.connect(
            host=HOST,
            user=USER,
            password=PASSWORD,
            database=DATABASE
        )

    except mysql.connector.Error as error:

        messagebox.showerror(
            "Error de conexión",
            f"No se pudo conectar con MySQL.\n\n{error}"
        )

        return None


def interfaz_administrador():

    ventana = tk.Tk()
    ventana.state("zoomed")
    ventana.title("Restaurante Vivaldi | Administrador")
    ventana.geometry("1100x720")
    ventana.configure(bg=CREMA)
    ventana.resizable(False, False)


    try:

        ventana.iconbitmap(LOGO_ICO)

    except:
        pass


    contenedor = tk.Frame(
        ventana,
        bg=CREMA
    )

    contenedor.pack(
        fill="both",
        expand=True
    )

    sidebar = tk.Frame(
        contenedor,
        bg=VERDE_OSCURO,
        width=230
    )

    sidebar.pack(
        side="left",
        fill="y"
    )

    sidebar.pack_propagate(False)

    contenido_sidebar = tk.Frame(
        sidebar,
        bg=VERDE_OSCURO
    )

    contenido_sidebar.place(
        relx=0.5,
        rely=0.5,
        anchor="center"
    )


    try:

        imagen_logo = Image.open(LOGO_PNG)
        imagen_logo = imagen_logo.resize((145, 145))

        logo = ImageTk.PhotoImage(imagen_logo)

        etiqueta_logo = tk.Label(
            contenido_sidebar,
            image=logo,
            bg=VERDE_OSCURO
        )

        etiqueta_logo.image = logo

        etiqueta_logo.pack(
            pady=(0, 12)
        )

    except:

        tk.Label(
            contenido_sidebar,
            text="VIVALDI",
            bg=VERDE_OSCURO,
            fg=BEIGE,
            font=("Georgia", 24, "bold")
        ).pack(
            pady=(0, 30)
        )

    tk.Label(
        contenido_sidebar,
        text="ADMINISTRACIÓN",
        bg=VERDE_OSCURO,
        fg=BEIGE,
        font=("Georgia", 13, "bold")
    ).pack(
        pady=(0, 20)
    )


    menu_navegacion = tk.Frame(
        contenido_sidebar,
        bg=VERDE_OSCURO
    )

    menu_navegacion.pack(
        fill="x",
        padx=15
    )

    contenido = tk.Frame(
        contenedor,
        bg=CREMA
    )

    contenido.pack(
        side="left",
        fill="both",
        expand=True
    )

    area_central = tk.Frame(
        contenido,
        bg=CREMA
    )

    area_central.place(
        relx=0.5,
        rely=0.5,
        anchor="center",
        relwidth=0.92,
        relheight=0.94
    )


    titulo_principal = tk.Label(
        area_central,
        text="Panel de Administración",
        bg=CREMA,
        fg=VERDE_OSCURO,
        font=("Georgia", 23, "bold")
    )

    titulo_principal.pack(
        anchor="w",
        pady=(5, 2)
    )


    subtitulo_principal = tk.Label(
        area_central,
        text="Supervisión y control general Restaurante Vivaldi",
        bg=CREMA,
        fg=GRIS,
        font=("Segoe UI", 10)
    )

    subtitulo_principal.pack(
        anchor="w",
        pady=(0, 20)
    )


    vista_actual = tk.Frame(
        area_central,
        bg=CREMA
    )

    vista_actual.pack(
        fill="both",
        expand=True
    )

    def mostrar_encabezados(mostrar):
        if mostrar:
            titulo_principal.pack(
                anchor="w",
                pady=(5, 2),
                before=vista_actual
            )

            subtitulo_principal.pack(
                anchor="w",
                pady=(0, 20),
                before=vista_actual
            )
        else:
            titulo_principal.pack_forget()
            subtitulo_principal.pack_forget()


    def limpiar_vista():
        for widget in vista_actual.winfo_children():

            widget.destroy()

    botones_menu = {}

    def activar_boton(nombre):

        for texto, boton in botones_menu.items():
            boton.config(
                text=f"   {texto}",
                bg=VERDE_OSCURO,
                fg=CREMA_CLARO,
                font=("Arial", 10, "bold")
            )

        if nombre in botones_menu:
            botones_menu[nombre].config(
                text=f"●  {nombre}",
                bg=VERDE,
                fg=BEIGE,
                font=("Arial", 10, "bold")
            )

    def crear_boton_menu(texto, comando):

        boton = tk.Button(
            menu_navegacion,
            text=f"   {texto}",
            bg=VERDE_OSCURO,
            fg=CREMA_CLARO,
            activebackground=VERDE,
            activeforeground=BEIGE,
            bd=0,
            relief="flat",
            anchor="w",
            padx=10,
            pady=13,
            font=("Arial", 10, "bold"),
            cursor="hand2",
            command=comando
        )

        boton.texto_menu = texto

        boton.pack(
            fill="x",
            pady=4
        )

        botones_menu[texto] = boton

        return boton

    def crear_tarjeta(padre, titulo, valor, descripcion):

        tarjeta = tk.Frame(
            padre,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        tarjeta.pack(
            side="left",
            fill="both",
            expand=True,
            padx=5
        )

        tk.Label(
            tarjeta,
            text=titulo,
            bg=BLANCO,
            fg=GRIS,
            font=("Segoe UI", 9, "bold")
        ).pack(
            anchor="w",
            padx=15,
            pady=(14, 2)
        )

        tk.Label(
            tarjeta,
            text=valor,
            bg=BLANCO,
            fg=VERDE_OSCURO,
            font=("Georgia", 22, "bold")
        ).pack(
            anchor="w",
            padx=15
        )

        tk.Label(
            tarjeta,
            text=descripcion,
            bg=BLANCO,
            fg=GRIS,
            font=("Segoe UI", 8)
        ).pack(
            anchor="w",
            padx=15,
            pady=(0, 14)
        )

        return tarjeta


    def mostrar_inicio():
        limpiar_vista()
        mostrar_encabezados(True)
        activar_boton("INICIO")

        marco_resumen = tk.Frame(vista_actual, bg=CREMA)
        marco_resumen.pack(fill="x", pady=(0, 15))

        ventas_hoy = 0.0
        total_mesas = 0
        platillos_disponibles = 0
        reservas_hoy = 0

        conexion = conectar()

        if conexion:
            cursor = None

            try:
                cursor = conexion.cursor()

                cursor.execute("""
                    SELECT COALESCE(SUM(total), 0)
                    FROM Factura
                    WHERE DATE(fecha) = CURDATE()
                    AND estado_pago = 'Pagado'
                """)
                ventas_hoy = float(cursor.fetchone()[0] or 0)

                cursor.execute("""
                    SELECT COUNT(*)
                    FROM Mesa
                """)
                total_mesas = cursor.fetchone()[0]

                cursor.execute("""
                    SELECT COUNT(*)
                    FROM Platillo
                    WHERE disponibilidad = 'Disponible'
                """)
                platillos_disponibles = cursor.fetchone()[0]

                cursor.execute("""
                    SELECT COUNT(*)
                    FROM Reserva
                    WHERE fecha = CURDATE()
                    AND estado <> 'Cancelada'
                """)
                reservas_hoy = cursor.fetchone()[0]

            except mysql.connector.Error as error:
                print("Error al cargar el resumen:", error)

            finally:
                if cursor:
                    cursor.close()

                conexion.close()

        crear_tarjeta(
            marco_resumen,
            "VENTAS DEL DÍA",
            f"${ventas_hoy:,.2f}",
            "Total de facturas pagadas hoy"
        )

        crear_tarjeta(
            marco_resumen,
            "PLATILLOS DISPONIBLES",
            platillos_disponibles,
            "Platillos disponibles en el menú"
        )

        crear_tarjeta(
            marco_resumen,
            "MESAS",
            total_mesas,
            "Mesas registradas"
        )

        crear_tarjeta(
            marco_resumen,
            "RESERVACIONES",
            reservas_hoy,
            "Reservaciones para hoy"
        )

        fila_inferior = tk.Frame(vista_actual, bg=CREMA)
        fila_inferior.pack(fill="both", expand=True)

        marco_reservaciones = tk.Frame(
            fila_inferior,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        marco_reservaciones.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(0, 7)
        )

        tk.Label(
            marco_reservaciones,
            text="PRÓXIMAS RESERVACIONES",
            bg=BLANCO,
            fg=VERDE_OSCURO,
            font=("Georgia", 13, "bold")
        ).pack(
            anchor="w",
            padx=18,
            pady=(16, 3)
        )

        tk.Label(
            marco_reservaciones,
            text="Reservaciones pendientes y confirmadas",
            bg=BLANCO,
            fg=GRIS,
            font=("Segoe UI", 8)
        ).pack(
            anchor="w",
            padx=18,
            pady=(0, 10)
        )

        canvas_reservaciones = tk.Canvas(
            marco_reservaciones,
            bg=BLANCO,
            highlightthickness=0
        )

        scrollbar_reservaciones = ttk.Scrollbar(
            marco_reservaciones,
            orient="vertical",
            command=canvas_reservaciones.yview
        )

        canvas_reservaciones.configure(
            yscrollcommand=scrollbar_reservaciones.set
        )

        scrollbar_reservaciones.pack(
            side="right",
            fill="y",
            padx=(0, 5),
            pady=5
        )

        canvas_reservaciones.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(10, 0),
            pady=(0, 10)
        )

        contenedor_reservaciones = tk.Frame(
            canvas_reservaciones,
            bg=BLANCO
        )

        ventana_reservaciones = canvas_reservaciones.create_window(
            (0, 0),
            window=contenedor_reservaciones,
            anchor="nw"
        )

        def actualizar_scroll_reservaciones(event=None):
            canvas_reservaciones.configure(
                scrollregion=canvas_reservaciones.bbox("all")
            )

        contenedor_reservaciones.bind(
            "<Configure>",
            actualizar_scroll_reservaciones
        )

        def ajustar_ancho_reservaciones(event):
            canvas_reservaciones.itemconfig(
                ventana_reservaciones,
                width=event.width
            )

        canvas_reservaciones.bind(
            "<Configure>",
            ajustar_ancho_reservaciones
        )

        def desplazar_reservaciones(event):
            canvas_reservaciones.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units"
            )

        canvas_reservaciones.bind(
            "<MouseWheel>",
            desplazar_reservaciones
        )

        contenedor_reservaciones.bind(
            "<MouseWheel>",
            desplazar_reservaciones
        )

        reserva_abierta = {
            "id": None,
            "frame": None
        }

        def mostrar_detalle_reserva(id_reserva, fecha,
                hora, cliente, personas, estado,
                mesa, contenedor_reserva):

            if reserva_abierta["id"] == id_reserva:
                if reserva_abierta["frame"] is not None:
                    reserva_abierta["frame"].destroy()

                reserva_abierta["frame"] = None
                reserva_abierta["id"] = None
                actualizar_scroll_reservaciones()
                return

            if reserva_abierta["frame"] is not None:
                reserva_abierta["frame"].destroy()

            reserva_abierta["id"] = id_reserva

            detalle = tk.Frame(
                contenedor_reserva,
                bg=CREMA,
                highlightbackground=BEIGE,
                highlightthickness=1
            )

            detalle.pack(
                fill="x",
                padx=18,
                pady=(0, 6)
            )

            reserva_abierta["frame"] = detalle

            tk.Label(
                detalle,
                text="DETALLES DE LA RESERVACIÓN",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w",
                padx=14,
                pady=(10, 7)
            )

            informacion = tk.Frame(
                detalle,
                bg=CREMA
            )

            informacion.pack(
                fill="x",
                padx=14,
                pady=(0, 10)
            )

            tk.Label(
                informacion,
                text=f"Cliente: {cliente or 'Sin cliente'}",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=2
            )

            tk.Label(
                informacion,
                text=f"Fecha: {fecha}",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=2
            )

            tk.Label(
                informacion,
                text=f"Hora: {str(hora)[:5]}",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=2
            )

            tk.Label(
                informacion,
                text=f"Personas: {personas}",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=2
            )

            tk.Label(
                informacion,
                text=f"Mesa: {mesa if mesa else 'Sin asignar'}",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(
                anchor="w",
                pady=2
            )

            tk.Label(
                informacion,
                text=f"Estado: {estado}",
                bg=CREMA,
                fg=VINO if estado == "Cancelada" else VERDE,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w",
                pady=2
            )

            actualizar_scroll_reservaciones()

        conexion = conectar()

        if conexion:
            cursor = None

            try:
                cursor = conexion.cursor()

                cursor.execute(
                    """
                    SELECT r.id_reserva, r.fecha, r.hora,
                        CONCAT(c.nombre, ' ', c.apellido) AS cliente,
                        r.cantidad_personas, r.estado,
                        m.numero_mesa
                    FROM Reserva r
                    LEFT JOIN Cliente c
                        ON r.id_cliente = c.id_cliente
                    LEFT JOIN Mesa m
                        ON r.id_mesa = m.id_mesa
                    WHERE
                        TIMESTAMP(r.fecha, r.hora) >= NOW()
                        AND r.estado <> 'Cancelada'
                    ORDER BY
                        r.fecha ASC,
                        r.hora ASC
                    LIMIT 5
                    """
                )

                reservas = cursor.fetchall()

                if reservas:

                    for reserva in reservas:

                        (
                            id_reserva,
                            fecha,
                            hora,
                            cliente,
                            personas,
                            estado,
                            mesa
                        ) = reserva

                        contenedor_reserva =tk.Frame(
                            contenedor_reservaciones,
                            bg=BLANCO
                        )

                        contenedor_reserva.pack(
                            fill="x",
                            pady=3
                        )

                        fila = tk.Frame(
                            contenedor_reserva,
                            bg=CREMA,
                            cursor="hand2"
                        )

                        fila.pack(
                            fill="x",
                            padx=18,
                            pady=3
                        )

                        if estado == "Confirmada":
                            color_estado = VERDE
                        elif estado == "Pendiente":
                            color_estado = BEIGE
                        else:
                            color_estado = GRIS

                        tk.Label(
                            fila,
                            text=str(hora)[:5],
                            bg=CREMA,
                            fg=VINO,
                            font=("Segoe UI", 9, "bold"),
                            width=8,
                            cursor="hand2"
                        ).pack(
                            side="left"
                        )

                        informacion_reserva = tk.Frame(
                            fila,
                            bg=CREMA,
                            cursor="hand2"
                        )

                        informacion_reserva.pack(
                            side="left",
                            fill="x",
                            expand=True
                        )

                        tk.Label(
                            informacion_reserva,
                            text=cliente or "Sin cliente",
                            bg=CREMA,
                            fg=VERDE_OSCURO,
                            font=("Segoe UI", 9, "bold"),
                            anchor="w",
                            cursor="hand2"
                        ).pack(
                            anchor="w"
                        )

                        tk.Label(
                            informacion_reserva,
                            text=f"{fecha}  •  {personas} personas",
                            bg=CREMA,
                            fg=GRIS,
                            font=("Segoe UI", 8),
                            anchor="w",
                            cursor="hand2"
                        ).pack(
                            anchor="w"
                        )

                        tk.Label(
                            fila,
                            text=estado,
                            bg=CREMA,
                            fg=color_estado,
                            font=("Segoe UI", 8, "bold"),
                            cursor="hand2"
                        ).pack(
                            side="right",
                            padx=(5, 0)
                        )

                        def seleccionar_reserva(event,
                                id_reserva=id_reserva,
                                fecha=fecha,
                                hora=hora,
                                cliente=cliente,
                                personas=personas,
                                estado=estado,
                                mesa=mesa,
                                contenedor_reserva=contenedor_reserva
                        ):
                            mostrar_detalle_reserva(
                                id_reserva,
                                fecha,
                                hora,
                                cliente,
                                personas,
                                estado,
                                mesa,
                                contenedor_reserva
                            )

                        fila.bind("<Button-1>", seleccionar_reserva)
                        informacion_reserva.bind(
                            "<Button-1>",
                            seleccionar_reserva
                        )

                        for widget in fila.winfo_children():
                            widget.bind(
                                "<Button-1>",
                                seleccionar_reserva
                            )

                else:

                    tk.Label(
                        contenedor_reservaciones,
                        text="No hay próximas reservaciones.",
                        bg=BLANCO,
                        fg=GRIS,
                        font=("Segoe UI", 9)
                    ).pack(
                        pady=30
                    )

            except mysql.connector.Error:

                tk.Label(
                    contenedor_reservaciones,
                    text="No se pudieron cargar las reservaciones.",
                    bg=BLANCO,
                    fg=VINO,
                    font=("Segoe UI", 9)
                ).pack(
                    pady=25
                )

            finally:

                if cursor:
                    cursor.close()

                conexion.close()

        marco_mesas = tk.Frame(
            fila_inferior,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        marco_mesas.pack(
            side="right",
            fill="both",
            expand=True,
            padx=(7, 0)
        )

        tk.Label(
            marco_mesas,
            text="ESTADO DE MESAS",
            bg=BLANCO,
            fg=VERDE_OSCURO,
            font=("Georgia", 13, "bold")
        ).pack(
            anchor="w",
            padx=18,
            pady=(16, 3)
        )

        tk.Label(
            marco_mesas,
            text="Situación actual del restaurante",
            bg=BLANCO,
            fg=GRIS,
            font=("Segoe UI", 8)
        ).pack(
            anchor="w",
            padx=18,
            pady=(0, 12)
        )

        contenedor_principal_mesas = tk.Frame(
            marco_mesas,
            bg=BLANCO
        )

        contenedor_principal_mesas.pack(
            fill="both",
            expand=True,
            padx=(12, 8),
            pady=(0, 10)
        )

        canvas_mesas_inicio = tk.Canvas(
            contenedor_principal_mesas,
            bg=BLANCO,
            highlightthickness=0
        )

        scrollbar_mesas_inicio = ttk.Scrollbar(
            contenedor_principal_mesas,
            orient="vertical",
            command=canvas_mesas_inicio.yview
        )

        canvas_mesas_inicio.configure(
            yscrollcommand=scrollbar_mesas_inicio.set
        )

        scrollbar_mesas_inicio.pack(
            side="right",
            fill="y"
        )

        canvas_mesas_inicio.pack(
            side="left",
            fill="both",
            expand=True
        )

        contenedor_mesas_inicio = tk.Frame(
            canvas_mesas_inicio,
            bg=BLANCO
        )

        ventana_mesas_inicio = canvas_mesas_inicio.create_window(
            (0, 0),
            window=contenedor_mesas_inicio,
            anchor="nw"
        )

        def actualizar_scroll_mesas_inicio(event=None):
            canvas_mesas_inicio.configure(
                scrollregion=canvas_mesas_inicio.bbox("all")
            )

        contenedor_mesas_inicio.bind(
            "<Configure>",
            actualizar_scroll_mesas_inicio
        )

        def ajustar_ancho_mesas_inicio(event):
            canvas_mesas_inicio.itemconfig(
                ventana_mesas_inicio,
                width=event.width
            )

        canvas_mesas_inicio.bind(
            "<Configure>",
            ajustar_ancho_mesas_inicio
        )

        def desplazar_mesas_inicio(event):
            canvas_mesas_inicio.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units"
            )

        canvas_mesas_inicio.bind(
            "<MouseWheel>",
            desplazar_mesas_inicio
        )

        contenedor_mesas_inicio.bind(
            "<MouseWheel>",
            desplazar_mesas_inicio
        )

        conexion = conectar()

        if conexion:

            cursor = None

            try:

                cursor = conexion.cursor()

                cursor.execute(
                    """
                    SELECT numero_mesa, estado
                    FROM Mesa
                    ORDER BY numero_mesa
                    """
                )

                mesas = cursor.fetchall()

                if mesas:

                    for indice, (numero, estado) in enumerate(mesas):

                        if estado == "Disponible":
                            color_estado = VERDE
                        elif estado == "Ocupada":
                            color_estado = VINO
                        else:
                            color_estado = BEIGE

                        mesa = tk.Frame(
                            contenedor_mesas_inicio,
                            bg=CREMA,
                            highlightbackground=GRIS_CLARO,
                            highlightthickness=1,
                            width=90,
                            height=60
                        )

                        mesa.grid(
                            row=indice // 3,
                            column=indice % 3,
                            padx=4,
                            pady=4,
                            sticky="nsew"
                        )

                        mesa.grid_propagate(False)

                        tk.Label(
                            mesa,
                            text=f"● Mesa {numero}",
                            bg=CREMA,
                            fg=color_estado,
                            font=("Segoe UI", 8, "bold")
                        ).place(
                            relx=0.5,
                            rely=0.5,
                            anchor="center"
                        )

                    for columna in range(3):
                        contenedor_mesas_inicio.grid_columnconfigure(
                            columna,
                            weight=1
                        )

                else:

                    tk.Label(
                        contenedor_mesas_inicio,
                        text="No hay mesas registradas.",
                        bg=BLANCO,
                        fg=GRIS,
                        font=("Segoe UI", 9)
                    ).pack(
                        pady=25
                    )

            except mysql.connector.Error:

                tk.Label(
                    contenedor_mesas_inicio,
                    text="No se pudieron cargar las mesas.",
                    bg=BLANCO,
                    fg=VINO,
                    font=("Segoe UI", 9)
                ).pack(
                    pady=25
                )

            finally:

                if cursor:
                    cursor.close()

                conexion.close()

    def mostrar_mesas():

        limpiar_vista()
        mostrar_encabezados(False)
        activar_boton("MESAS")

        encabezado = tk.Frame(
            vista_actual,
            bg=CREMA
        )

        encabezado.pack(
            fill="x",
            pady=(10, 15)
        )

        tk.Label(
            encabezado,
            text="Mesas",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Georgia", 22, "bold")
        ).pack(
            anchor="w",
        )

        tk.Label(
            encabezado,
            text="Administración y estado de las mesas",
            bg=CREMA,
            fg=GRIS,
            font=("Segoe UI", 9)
        ).pack(
            anchor="w",
            pady=(8, 0)
        )

        barra_busqueda = tk.Frame(
            vista_actual,
            bg=CREMA
        )

        barra_busqueda.pack(
            fill="x",
            pady=(0, 12)
        )

        tk.Label(
            barra_busqueda,
            text="Buscar mesa:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(
            side="left"
        )

        entrada_busqueda = tk.Entry(
            barra_busqueda,
            font=("Segoe UI", 10),
            bg=BLANCO,
            fg=VERDE_OSCURO,
            relief="flat",
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        entrada_busqueda.pack(
            side="left",
            padx=(8, 10),
            ipady=6,
            fill="x",
            expand=True
        )

        marco_mesas = tk.Frame(
            vista_actual,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        marco_mesas.pack(
            fill="both",
            expand=True
        )

        canvas_mesas = tk.Canvas(
            marco_mesas,
            bg=BLANCO,
            highlightthickness=0
        )

        scrollbar_mesas = ttk.Scrollbar(
            marco_mesas,
            orient="vertical",
            command=canvas_mesas.yview
        )

        canvas_mesas.configure(
            yscrollcommand=scrollbar_mesas.set
        )

        scrollbar_mesas.pack(
            side="right",
            fill="y"
        )

        canvas_mesas.pack(
            side="left",
            fill="both",
            expand=True
        )

        def desplazar_mesas(event):
            canvas_mesas.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units"
            )

        canvas_mesas.bind(
            "<MouseWheel>",
            desplazar_mesas
        )

        contenedor_mesas = tk.Frame(
            canvas_mesas,
            bg=BLANCO
        )

        ventana_mesas = canvas_mesas.create_window(
            (0, 0),
            window=contenedor_mesas,
            anchor="nw"
        )

        def actualizar_scroll(event=None):

            canvas_mesas.configure(
                scrollregion=canvas_mesas.bbox("all")
            )

        contenedor_mesas.bind(
            "<Configure>",
            actualizar_scroll
        )

        def ajustar_ancho(event):

            canvas_mesas.itemconfig(
                ventana_mesas,
                width=event.width
            )

        canvas_mesas.bind(
            "<Configure>",
            ajustar_ancho
        )

        def cargar_mesas():

            for widget in contenedor_mesas.winfo_children():
                widget.destroy()

            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:

                cursor = conexion.cursor()

                texto = entrada_busqueda.get().strip()

                consulta = """
                    SELECT
                        id_mesa,
                        numero_mesa,
                        capacidad,
                        zona,
                        estado
                    FROM Mesa
                """

                parametros = ()

                if texto:
                    consulta += """
                        WHERE
                            CAST(numero_mesa AS CHAR) LIKE %s
                            OR zona LIKE %s
                            OR estado LIKE %s
                    """

                    buscar = f"%{texto}%"

                    parametros = (
                        buscar,
                        buscar,
                        buscar
                    )

                consulta += """
                    ORDER BY numero_mesa
                """

                cursor.execute(
                    consulta,
                    parametros
                )

                mesas = cursor.fetchall()

                if not mesas:
                    tk.Label(
                        contenedor_mesas,
                        text="No se encontraron mesas.",
                        bg=BLANCO,
                        fg=GRIS,
                        font=("Segoe UI", 10)
                    ).pack(
                        pady=40
                    )

                    return

                for indice, mesa in enumerate(mesas):

                    id_mesa = mesa[0]
                    numero = mesa[1]
                    capacidad = mesa[2]
                    zona = mesa[3]
                    estado = mesa[4]

                    if estado == "Disponible":

                        color_estado = VERDE

                    elif estado == "Ocupada":

                        color_estado = VINO

                    else:

                        color_estado = BEIGE

                    tarjeta = tk.Frame(
                        contenedor_mesas,
                        bg=CREMA,
                        highlightbackground=GRIS_CLARO,
                        highlightthickness=3,
                        width=180,
                        height=145
                    )

                    tarjeta.grid(
                        row=indice // 4,
                        column=indice % 4,
                        padx=8,
                        pady=8,
                        sticky="nsew"
                    )

                    tarjeta.grid_propagate(False)

                    tk.Label(
                        tarjeta,
                        text=f"MESA {numero}",
                        bg=CREMA,
                        fg=VERDE_OSCURO,
                        font=("Georgia", 14, "bold")
                    ).pack(
                        pady=(15, 3)
                    )

                    tk.Label(
                        tarjeta,
                        text=f"Capacidad: {capacidad} personas",
                        bg=CREMA,
                        fg=GRIS,
                        font=("Segoe UI", 9)
                    ).pack()

                    tk.Label(
                        tarjeta,
                        text=f"Zona: {zona}",
                        bg=CREMA,
                        fg=GRIS,
                        font=("Segoe UI", 9)
                    ).pack(
                        pady=(3, 5)
                    )

                    tk.Label(
                        tarjeta,
                        text=f"● {estado.upper()}",
                        bg=CREMA,
                        fg=color_estado,
                        font=("Segoe UI", 9, "bold")
                    ).pack()

                    tarjeta.bind(
                        "<Button-1>",
                        lambda event, datos=mesa, marco=tarjeta: seleccionar_mesa(datos, marco)
                    )

                    for elemento in tarjeta.winfo_children():
                        elemento.bind(
                            "<Button-1>",
                            lambda event, datos=mesa, marco=tarjeta: seleccionar_mesa(datos, marco)
                        )

                for columna in range(4):
                    contenedor_mesas.grid_columnconfigure(
                        columna,
                        weight=1
                    )


            except mysql.connector.Error as error:

                messagebox.showerror(
                    "Error",
                    f"No se pudieron cargar las mesas.\n\n{error}"
                )


            finally:

                if cursor:
                    cursor.close()

                conexion.close()

        def abrir_formulario_mesa(mesa=None):

            formulario = tk.Toplevel(
                ventana
            )

            formulario.title(
                "Nueva mesa"
                if mesa is None
                else "Editar mesa"
            )

            formulario.geometry(
                "420x450"
            )

            formulario.configure(
                bg=CREMA
            )

            formulario.resizable(
                False,
                False
            )

            try:

                formulario.iconbitmap(
                    LOGO_ICO
                )

            except:

                pass

            formulario.transient(
                ventana
            )

            formulario.grab_set()

            tk.Label(
                formulario,
                text=(
                    "Nueva mesa"
                    if mesa is None
                    else "Editar mesa"
                ),
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Georgia", 19, "bold")
            ).pack(
                pady=(25, 5)
            )

            tk.Label(
                formulario,
                text="Información de la mesa",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(
                pady=(0, 20)
            )

            campos = tk.Frame(
                formulario,
                bg=CREMA
            )

            campos.pack(
                padx=40,
                fill="x"
            )

            tk.Label(
                campos,
                text="Número de mesa",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w"
            )

            entrada_numero = tk.Entry(
                campos,
                font=("Segoe UI", 10),
                bg=BLANCO,
                fg=VERDE_OSCURO,
                relief="flat",
                highlightbackground=GRIS_CLARO,
                highlightthickness=1
            )

            entrada_numero.pack(
                fill="x",
                ipady=6,
                pady=(3, 10)
            )

            tk.Label(
                campos,
                text="Capacidad",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w"
            )

            entrada_capacidad = tk.Entry(
                campos,
                font=("Segoe UI", 10),
                bg=BLANCO,
                fg=VERDE_OSCURO,
                relief="flat",
                highlightbackground=GRIS_CLARO,
                highlightthickness=1
            )

            entrada_capacidad.pack(
                fill="x",
                ipady=6,
                pady=(3, 10)
            )

            tk.Label(
                campos,
                text="Zona",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w"
            )

            zona = ttk.Combobox(
                campos,
                values=[
                    "Salón",
                    "Terraza",
                    "Jardín"
                ],
                state="readonly",
                font=("Segoe UI", 10)
            )

            zona.pack(
                fill="x",
                pady=(3, 10)
            )

            tk.Label(
                campos,
                text="Estado",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(
                anchor="w"
            )

            estado = ttk.Combobox(
                campos,
                values=[
                    "Disponible",
                    "Ocupada",
                    "Reservada"
                ],
                state="readonly",
                font=("Segoe UI", 10)
            )

            estado.pack(
                fill="x",
                pady=(3, 10)
            )

            if mesa:

                entrada_numero.insert(
                    0,
                    mesa[1]
                )

                entrada_capacidad.insert(
                    0,
                    mesa[2]
                )

                zona.set(
                    mesa[3]
                )

                estado.set(
                    mesa[4]
                )

            else:

                estado.set(
                    "Disponible"
                )

            def guardar_mesa():

                numero = entrada_numero.get().strip()
                capacidad = entrada_capacidad.get().strip()
                zona_valor = zona.get()
                estado_valor = estado.get()

                if not numero or not capacidad or not zona_valor:
                    messagebox.showwarning(
                        "Datos incompletos",
                        "Completa todos los campos.",
                        parent=formulario
                    )

                    return

                try:

                    numero = int(numero)
                    capacidad = int(capacidad)

                    if numero <= 0 or capacidad <= 0:
                        raise ValueError


                except ValueError:

                    messagebox.showwarning(
                        "Datos inválidos",
                        "El número de mesa y la capacidad deben ser números positivos.",
                        parent=formulario
                    )

                    return

                conexion = conectar()

                if not conexion:
                    return

                cursor = None

                try:

                    cursor = conexion.cursor()

                    if mesa is None:

                        cursor.execute(
                            """
                            INSERT INTO Mesa
                            (numero_mesa, capacidad,
                                zona, estado)
                            VALUES(%s, %s, %s, %s)
                            """,
                            (numero, capacidad,
                                zona_valor, estado_valor)
                        )

                        mensaje = "Mesa registrada correctamente."

                    else:

                        cursor.execute(
                            """
                            UPDATE Mesa
                            SET
                                numero_mesa = %s, capacidad = %s,
                                zona = %s, estado = %s
                            WHERE id_mesa = %s
                            """,
                            (numero, capacidad, zona_valor,
                                estado_valor, mesa[0])
                        )

                        mensaje = "Mesa actualizada correctamente."

                    conexion.commit()
                    formulario.destroy()
                    cargar_mesas()

                    messagebox.showinfo(
                        "Operación realizada",
                        mensaje
                    )

                except mysql.connector.Error as error:

                    conexion.rollback()

                    messagebox.showerror(
                        "Error",
                        f"No se pudo guardar la mesa.\n\n{error}",
                        parent=formulario
                    )

                finally:

                    if cursor:
                        cursor.close()

                    conexion.close()

            tk.Button(
                formulario,
                text="Guardar",
                command=guardar_mesa,
                bg=VERDE,
                fg=BEIGE,
                activebackground=VERDE_OSCURO,
                activeforeground=BEIGE,
                relief="flat",
                bd=0,
                font=("Segoe UI", 9, "bold"),
                padx=25,
                pady=9,
                cursor="hand2"
            ).pack(
                side="left",
                padx=(40, 5),
                pady=15
            )

            tk.Button(
                formulario,
                text="Cancelar",
                command=formulario.destroy,
                bg=GRIS_CLARO,
                fg=VERDE_OSCURO,
                activebackground=GRIS_CLARO,
                activeforeground=VERDE_OSCURO,
                relief="flat",
                bd=0,
                font=("Segoe UI", 9, "bold"),
                padx=25,
                pady=9,
                cursor="hand2"
            ).pack(
                side="left",
                padx=5,
                pady=15
            )

        mesa_seleccionada = [None]

        def seleccionar_mesa(mesa, tarjeta):
            mesa_seleccionada[0] = mesa
            for widget in contenedor_mesas.winfo_children():

                if isinstance(widget, tk.Frame):
                    widget.config(
                        highlightbackground=GRIS_CLARO,
                        highlightthickness=2
                    )

            tarjeta.config(
                highlightbackground=BEIGE,
                highlightthickness=2
            )

        def editar_mesa():

            if mesa_seleccionada[0] is None:
                messagebox.showwarning(
                    "Seleccionar mesa",
                    "Selecciona una mesa para editar."
                )

                return

            abrir_formulario_mesa(
                mesa_seleccionada[0]
            )

        def eliminar_mesa():

            if mesa_seleccionada[0] is None:
                messagebox.showwarning(
                    "Seleccionar mesa",
                    "Selecciona una mesa para eliminar."
                )
                return

            mesa = mesa_seleccionada[0]

            confirmar = messagebox.askyesno(
                "Confirmar eliminación",
                f"¿Deseas eliminar la mesa {mesa[1]}?"
            )

            if not confirmar:
                return

            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:
                cursor = conexion.cursor()
                cursor.execute(
                    """
                    DELETE FROM Mesa
                    WHERE id_mesa = %s
                    """,
                    (mesa[0],)
                )

                conexion.commit()
                mesa_seleccionada[0] = None
                cargar_mesas()

                messagebox.showinfo(
                    "Mesa eliminada",
                    "La mesa fue eliminada correctamente."
                )

            except mysql.connector.Error:

                conexion.rollback()
                messagebox.showerror(
                    "No se puede eliminar",
                    "No se puede eliminar esta mesa porque tiene "
                    "reservaciones relacionadas."
                )

            finally:

                if cursor:
                    cursor.close()

                conexion.close()

        barra_botones = tk.Frame(
            vista_actual,
            bg=CREMA
        )

        barra_botones.pack(
            fill="x",
            pady=(12, 0)
        )

        tk.Button(
            barra_botones,
            text="+ Nueva mesa",
            command=lambda: abrir_formulario_mesa(),
            bg=VERDE,
            fg=BEIGE,
            activebackground=VERDE_OSCURO,
            activeforeground=BEIGE,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            width=14,
            height=2,
            cursor="hand2"
        ).pack(
            side="left"
        )

        tk.Button(
            barra_botones,
            text="Editar",
            command=editar_mesa,
            bg=BEIGE,
            fg=VERDE_OSCURO,
            activebackground=BEIGE,
            activeforeground=VERDE_OSCURO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            width=14,
            height=2,
            cursor="hand2"
        ).pack(
            side="left",
            padx=8
        )

        tk.Button(
            barra_botones,
            text="Eliminar",
            command=eliminar_mesa,
            bg=VINO,
            fg=BLANCO,
            activebackground=VINO,
            activeforeground=BLANCO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            width=14,
            height=2,
            cursor="hand2"
        ).pack(
            side="left"
        )

        entrada_busqueda.bind(
            "<KeyRelease>",
            lambda event: cargar_mesas()
        )
        cargar_mesas()

    def mostrar_inventario():

        limpiar_vista()
        mostrar_encabezados(False)
        activar_boton("INVENTARIO")

        encabezado = tk.Frame(vista_actual, bg=CREMA)
        encabezado.pack(fill="x", pady=(8, 5))

        tk.Label(
            encabezado,
            text="Inventario",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Georgia", 23, "bold")
        ).pack(anchor="w")

        tk.Label(
            encabezado,
            text="Control de productos, ingredientes y existencias del restaurante",
            bg=CREMA,
            fg=GRIS,
            font=("Segoe UI", 9)
        ).pack(anchor="w", pady=(4, 0))

        marco_resumen = tk.Frame(vista_actual, bg=CREMA)
        marco_resumen.pack(fill="x", pady=(18, 15))

        def crear_tarjeta_inventario(padre, titulo, descripcion, color=VERDE_OSCURO):

            tarjeta = tk.Frame(
                padre,
                bg=BLANCO,
                highlightbackground=GRIS_CLARO,
                highlightthickness=1
            )

            tarjeta.pack(
                side="left",
                fill="both",
                expand=True,
                padx=(0, 10)
            )

            tk.Label(
                tarjeta,
                text=titulo,
                bg=BLANCO,
                fg=GRIS,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=15, pady=(14, 2))

            etiqueta_valor = tk.Label(
                tarjeta,
                text="0",
                bg=BLANCO,
                fg=color,
                font=("Georgia", 21, "bold")
            )
            etiqueta_valor.pack(anchor="w", padx=15)

            tk.Label(
                tarjeta,
                text=descripcion,
                bg=BLANCO,
                fg=GRIS,
                font=("Segoe UI", 8)
            ).pack(anchor="w", padx=15, pady=(2, 14))

            return etiqueta_valor

        valor_productos = crear_tarjeta_inventario(
            marco_resumen,
            "PRODUCTOS",
            "Ingredientes registrados"
        )

        valor_disponibles = crear_tarjeta_inventario(
            marco_resumen,
            "DISPONIBLES",
            "Productos con existencia"
        )

        valor_stock_bajo = crear_tarjeta_inventario(
            marco_resumen,
            "STOCK BAJO",
            "Productos por reponer"
        )

        valor_agotados = crear_tarjeta_inventario(
            marco_resumen,
            "AGOTADOS",
            "Productos sin existencia",
            VINO
        )

        barra_busqueda = tk.Frame(vista_actual, bg=CREMA)
        barra_busqueda.pack(fill="x", pady=(3, 12))

        tk.Label(
            barra_busqueda,
            text="Buscar:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left")

        entrada_busqueda = tk.Entry(
            barra_busqueda,
            font=("Segoe UI", 10),
            bg=BLANCO,
            fg=VERDE_OSCURO,
            relief="flat",
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        entrada_busqueda.pack(
            side="left",
            padx=(8, 15),
            ipady=6,
            fill="x",
            expand=True
        )

        tk.Label(
            barra_busqueda,
            text="Categoría:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 7))

        filtro_categoria = ttk.Combobox(
            barra_busqueda,
            values=["Todas"],
            state="readonly",
            width=16,
            font=("Segoe UI", 9)
        )
        filtro_categoria.set("Todas")
        filtro_categoria.pack(side="left", padx=(0, 10))

        tk.Label(
            barra_busqueda,
            text="Estado:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 7))

        filtro_estado = ttk.Combobox(
            barra_busqueda,
            values=["Todos", "Disponible", "Stock bajo", "Agotado"],
            state="readonly",
            width=15,
            font=("Segoe UI", 9)
        )
        filtro_estado.set("Todos")
        filtro_estado.pack(side="left")

        barra_acciones = tk.Frame(vista_actual, bg=CREMA)
        barra_acciones.pack(fill="x", pady=(0, 12))

        contenedor_inventario = tk.Frame(
            vista_actual,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )
        contenedor_inventario.pack(fill="both", expand=True)

        encabezado_tabla = tk.Frame(
            contenedor_inventario,
            bg=CREMA_CLARO
        )
        encabezado_tabla.pack(fill="x")

        columnas = [
            ("PRODUCTO", 25),
            ("CATEGORÍA", 18),
            ("CANTIDAD", 12),
            ("UNIDAD", 12),
            ("STOCK MÍNIMO", 15),
            ("ESTADO", 15)
        ]

        for texto_columna, ancho in columnas:
            tk.Label(
                encabezado_tabla,
                text=texto_columna,
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 8, "bold"),
                anchor="w",
                width=ancho
            ).pack(side="left", padx=8, pady=11)

        area_filas = tk.Frame(contenedor_inventario, bg=BLANCO)
        area_filas.pack(fill="both", expand=True)

        lienzo = tk.Canvas(
            area_filas,
            bg=BLANCO,
            highlightthickness=0
        )

        scroll_inventario = ttk.Scrollbar(
            area_filas,
            orient="vertical",
            command=lienzo.yview
        )

        lienzo.configure(yscrollcommand=scroll_inventario.set)

        scroll_inventario.pack(side="right", fill="y")
        lienzo.pack(side="left", fill="both", expand=True)

        marco_filas = tk.Frame(lienzo, bg=BLANCO)

        ventana_filas = lienzo.create_window(
            (0, 0),
            window=marco_filas,
            anchor="nw"
        )

        marco_filas.bind(
            "<Configure>",
            lambda evento: lienzo.configure(
                scrollregion=lienzo.bbox("all")
            )
        )

        lienzo.bind(
            "<Configure>",
            lambda evento: lienzo.itemconfigure(
                ventana_filas,
                width=evento.width
            )
        )

        ingredientes = []

        def obtener_ingredientes():
            conexion = None
            cursor = None

            try:
                conexion = conectar()

                if conexion is None:
                    return []

                cursor = conexion.cursor(dictionary=True)

                cursor.execute("""
                    SELECT id_ingrediente, nombre, categoria,
                           cantidad_actual, unidad_medida,
                           stock_minimo, costo_unitario,
                           id_proveedor, activo
                    FROM ingrediente
                    WHERE activo = 1
                    ORDER BY nombre ASC
                """)

                return cursor.fetchall()

            except mysql.connector.Error as error:
                messagebox.showerror(
                    "Error de inventario",
                    f"No se pudieron consultar los ingredientes.\n\n{error}"
                )
                return []

            finally:
                if cursor is not None:
                    cursor.close()

                if conexion is not None:
                    conexion.close()

        def determinar_estado(ingrediente):

            cantidad = float(ingrediente["cantidad_actual"] or 0)
            minimo = float(ingrediente["stock_minimo"] or 0)

            if cantidad <= 0:
                return "Agotado"

            if cantidad <= minimo:
                return "Stock bajo"

            return "Disponible"

        def actualizar_tabla(*args):

            for widget in marco_filas.winfo_children():
                widget.destroy()

            busqueda = entrada_busqueda.get().strip().lower()
            categoria_elegida = filtro_categoria.get()
            estado_elegido = filtro_estado.get()

            encontrados = 0

            for ingrediente in ingredientes:

                nombre = str(ingrediente["nombre"] or "")
                categoria = str(
                    ingrediente["categoria"] or "Sin categoría"
                )
                estado = determinar_estado(ingrediente)

                if busqueda and busqueda not in nombre.lower():
                    continue

                if (
                        categoria_elegida != "Todas"
                        and categoria != categoria_elegida
                ):
                    continue

                if estado_elegido != "Todos" and estado != estado_elegido:
                    continue

                color_fila = BLANCO if encontrados % 2 == 0 else CREMA_CLARO

                fila = tk.Frame(
                    marco_filas,
                    bg=color_fila,
                    cursor="hand2"
                )
                fila.pack(fill="x")

                datos = [
                    nombre,
                    categoria,
                    f'{float(ingrediente["cantidad_actual"] or 0):g}',
                    ingrediente["unidad_medida"],
                    f'{float(ingrediente["stock_minimo"] or 0):g}',
                    estado
                ]

                for indice, dato in enumerate(datos):
                    color_texto = (
                        VINO if dato == "Agotado"
                        else "#A66B16" if dato == "Stock bajo"
                        else VERDE_OSCURO
                    )

                    etiqueta = tk.Label(
                        fila,
                        text=str(dato),
                        bg=color_fila,
                        fg=color_texto,
                        font=("Segoe UI", 9),
                        anchor="w",
                        width=columnas[indice][1]
                    )
                    etiqueta.pack(side="left", padx=8, pady=12)

                    etiqueta.bind(
                        "<Double-Button-1>",
                        lambda evento, item=ingrediente: abrir_formulario(item)
                    )

                fila.bind(
                    "<Double-Button-1>",
                    lambda evento, item=ingrediente: abrir_formulario(item)
                )

                tk.Frame(
                    marco_filas,
                    bg=GRIS_CLARO,
                    height=1
                ).pack(fill="x")

                encontrados += 1

            if encontrados == 0:
                tk.Label(
                    marco_filas,
                    text="No hay productos que coincidan con la búsqueda.",
                    bg=BLANCO,
                    fg=GRIS,
                    font=("Segoe UI", 10)
                ).pack(pady=45)

            lienzo.configure(scrollregion=lienzo.bbox("all"))

        def cargar_inventario():

            nonlocal ingredientes

            ingredientes = obtener_ingredientes()

            valor_productos.configure(text=str(len(ingredientes)))

            valor_disponibles.configure(
                text=str(sum(
                    1 for item in ingredientes
                    if determinar_estado(item) == "Disponible"
                ))
            )

            valor_stock_bajo.configure(
                text=str(sum(
                    1 for item in ingredientes
                    if determinar_estado(item) == "Stock bajo"
                ))
            )

            valor_agotados.configure(
                text=str(sum(
                    1 for item in ingredientes
                    if determinar_estado(item) == "Agotado"
                ))
            )

            categorias = sorted({
                str(item["categoria"])
                for item in ingredientes
                if item["categoria"]
            })

            categoria_anterior = filtro_categoria.get()

            filtro_categoria["values"] = ["Todas"] + categorias

            if categoria_anterior in categorias:
                filtro_categoria.set(categoria_anterior)
            else:
                filtro_categoria.set("Todas")

            actualizar_tabla()

        def abrir_formulario(ingrediente=None):

            ventana = tk.Toplevel(vista_actual)
            ventana.title(
                "Editar ingrediente" if ingrediente else "Nuevo producto"
            )
            ventana.geometry("430x530")
            ventana.configure(bg=CREMA)
            ventana.resizable(False, False)
            ventana.transient(vista_actual.winfo_toplevel())
            ventana.grab_set()
            ventana.iconbitmap("vivaldi.ico")

            tk.Label(
                ventana,
                text="Editar ingrediente" if ingrediente else "Nuevo producto",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Georgia", 18, "bold")
            ).pack(anchor="w", padx=25, pady=(20, 15))

            campos = {}

            definiciones = [
                ("nombre", "Nombre del ingrediente"),
                ("categoria", "Categoría"),
                ("cantidad_actual", "Cantidad actual"),
                ("unidad_medida", "Unidad de medida"),
                ("stock_minimo", "Stock mínimo"),
                ("costo_unitario", "Costo unitario")
            ]

            for campo, titulo in definiciones:

                tk.Label(
                    ventana,
                    text=titulo,
                    bg=CREMA,
                    fg=VERDE_OSCURO,
                    font=("Segoe UI", 9, "bold")
                ).pack(anchor="w", padx=25, pady=(5, 2))

                entrada = tk.Entry(
                    ventana,
                    font=("Segoe UI", 10),
                    relief="flat",
                    highlightbackground=BEIGE,
                    highlightthickness=1
                )
                entrada.pack(fill="x", padx=25, ipady=5)

                if ingrediente:
                    valor = ingrediente.get(campo, "")
                    entrada.insert(
                        0,
                        str(valor if valor is not None else "")
                    )

                campos[campo] = entrada

            def guardar():

                nombre = campos["nombre"].get().strip()
                categoria = campos["categoria"].get().strip() or None
                unidad = campos["unidad_medida"].get().strip()

                if not nombre or not unidad:
                    messagebox.showwarning(
                        "Datos incompletos",
                        "El nombre y la unidad de medida son obligatorios.",
                        parent=ventana
                    )
                    return

                try:
                    cantidad = float(
                        campos["cantidad_actual"].get() or 0
                    )
                    minimo = float(
                        campos["stock_minimo"].get() or 0
                    )
                    costo = float(
                        campos["costo_unitario"].get() or 0
                    )

                    if cantidad < 0 or minimo < 0 or costo < 0:
                        raise ValueError("No se permiten valores negativos.")

                except ValueError as error:
                    messagebox.showerror(
                        "Datos incorrectos",
                        f"Revisa las cantidades y el costo.\n\n{error}",
                        parent=ventana
                    )
                    return

                conexion = None
                cursor = None

                try:
                    conexion = conectar()

                    if conexion is None:
                        return

                    cursor = conexion.cursor()

                    datos = (nombre, categoria, cantidad,
                        unidad, minimo, costo
                    )

                    if ingrediente:

                        cursor.execute("""
                            UPDATE ingrediente
                            SET nombre = %s,
                                categoria = %s,
                                cantidad_actual = %s,
                                unidad_medida = %s,
                                stock_minimo = %s,
                                costo_unitario = %s
                            WHERE id_ingrediente = %s
                        """, datos + (ingrediente["id_ingrediente"],))

                    else:

                        cursor.execute("""
                            INSERT INTO ingrediente (
                                nombre,
                                categoria,
                                cantidad_actual,
                                unidad_medida,
                                stock_minimo,
                                costo_unitario,
                                id_proveedor,
                                activo
                            )
                            VALUES (%s, %s, %s, %s, %s, %s, NULL, 1)
                        """, datos)

                    conexion.commit()
                    ventana.destroy()
                    cargar_inventario()

                    messagebox.showinfo(
                        "Inventario",
                        "Ingrediente guardado correctamente.",
                        parent=vista_actual.winfo_toplevel()
                    )

                except Exception as error:

                    if conexion is not None:
                        conexion.rollback()

                    messagebox.showerror(
                        "Error",
                        f"No se pudo guardar el ingrediente.\n\n{error}",
                        parent=ventana
                    )

                finally:

                    if cursor is not None:
                        cursor.close()

                    if conexion is not None:
                        conexion.close()

            tk.Button(
                ventana,
                text="Guardar ingrediente",
                command=guardar,
                bg=VERDE,
                fg=BEIGE,
                activebackground=VERDE_OSCURO,
                activeforeground=BEIGE,
                relief="flat",
                font=("Segoe UI", 10, "bold"),
                cursor="hand2",
                pady=9
            ).pack(fill="x", padx=25, pady=20)

        tk.Button(
            barra_acciones,
            text="+ Nuevo producto",
            command=lambda: abrir_formulario(),
            bg=VERDE,
            fg=BEIGE,
            activebackground=VERDE_OSCURO,
            activeforeground=BEIGE,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2",
            padx=18,
            pady=9
        ).pack(side="left")

        tk.Button(
            barra_acciones,
            text="Actualizar",
            command=cargar_inventario,
            bg=BEIGE,
            fg=VERDE_OSCURO,
            activebackground=BEIGE,
            activeforeground=VERDE_OSCURO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2",
            padx=18,
            pady=9
        ).pack(side="left", padx=8)

        entrada_busqueda.bind("<KeyRelease>", actualizar_tabla)
        filtro_categoria.bind("<<ComboboxSelected>>", actualizar_tabla)
        filtro_estado.bind("<<ComboboxSelected>>", actualizar_tabla)

        cargar_inventario()

    def mostrar_proveedores():
        limpiar_vista()
        mostrar_encabezados(False)
        activar_boton("PROVEEDORES")

        encabezado = tk.Frame(vista_actual, bg=CREMA)
        encabezado.pack(fill="x", pady=(8, 12))

        tk.Label(
            encabezado,
            text="Gestión de proveedores",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Georgia", 22, "bold")
        ).pack(anchor="w")

        tk.Label(
            encabezado,
            text="Administra los contactos y la información de tus proveedores",
            bg=CREMA,
            fg=GRIS,
            font=("Segoe UI", 9)
        ).pack(anchor="w", pady=(4, 0))

        marco_resumen = tk.Frame(vista_actual, bg=CREMA)
        marco_resumen.pack(fill="x", pady=(5, 15))

        def crear_tarjeta_proveedor(padre, titulo, color):
            tarjeta = tk.Frame(
                padre,
                bg=BLANCO,
                highlightbackground=GRIS_CLARO,
                highlightthickness=1
            )

            tarjeta.pack(
                side="left",
                fill="both",
                expand=True,
                padx=(0, 8)
            )

            tk.Label(
                tarjeta,
                text=titulo,
                bg=BLANCO,
                fg=GRIS,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=13, pady=(10, 2))

            valor = tk.Label(
                tarjeta,
                text="0",
                bg=BLANCO,
                fg=color,
                font=("Georgia", 20, "bold")
            )

            valor.pack(anchor="w", padx=13, pady=(0, 10))
            return valor

        total_proveedores = crear_tarjeta_proveedor(
            marco_resumen, "TOTAL DE PROVEEDORES", VERDE_OSCURO
        )

        proveedores_activos = crear_tarjeta_proveedor(
            marco_resumen, "ACTIVOS", VERDE
        )

        proveedores_inactivos = crear_tarjeta_proveedor(
            marco_resumen, "INACTIVOS", VINO
        )

        barra_busqueda = tk.Frame(vista_actual, bg=CREMA)
        barra_busqueda.pack(fill="x", pady=(0, 10))

        tk.Label(
            barra_busqueda,
            text="Buscar:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left")

        entrada_busqueda = tk.Entry(
            barra_busqueda,
            bg=BLANCO,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 10),
            relief="flat",
            highlightbackground=BEIGE,
            highlightthickness=1
        )

        entrada_busqueda.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(8, 12),
            ipady=6
        )

        tk.Label(
            barra_busqueda,
            text="Estado:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 7))

        filtro_estado = ttk.Combobox(
            barra_busqueda,
            values=["Todos", "Activos", "Inactivos"],
            state="readonly",
            width=12,
            font=("Segoe UI", 9)
        )

        filtro_estado.set("Activos")
        filtro_estado.pack(side="left")

        barra_acciones = tk.Frame(vista_actual, bg=CREMA)
        barra_acciones.pack(
            side="top",
            fill="x",
            pady=(0, 10)
        )

        marco_tabla = tk.Frame(
            vista_actual,
            bg=BLANCO,
            highlightbackground=GRIS_CLARO,
            highlightthickness=1
        )

        marco_tabla.pack(fill="both", expand=True)

        estilo_proveedores = ttk.Style()

        estilo_proveedores.configure(
            "Proveedores.Treeview",
            background=BLANCO,
            foreground=VERDE_OSCURO,
            fieldbackground=BLANCO,
            rowheight=34,
            font=("Segoe UI", 9)
        )

        estilo_proveedores.configure(
            "Proveedores.Treeview.Heading",
            background=CREMA_CLARO,
            foreground=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold"),
            padding=9
        )

        estilo_proveedores.map(
            "Proveedores.Treeview",
            background=[("selected", VERDE)],
            foreground=[("selected", BLANCO)]
        )

        columnas = (
            "id",
            "nombre",
            "telefono",
            "correo",
            "direccion",
            "estado"
        )

        tabla = ttk.Treeview(
            marco_tabla,
            columns=columnas,
            show="headings",
            style="Proveedores.Treeview",
            selectmode="browse"
        )

        tabla.heading("id", text="ID")
        tabla.heading("nombre", text="PROVEEDOR")
        tabla.heading("telefono", text="TELÉFONO")
        tabla.heading("correo", text="CORREO")
        tabla.heading("direccion", text="DIRECCIÓN")
        tabla.heading("estado", text="ESTADO")

        tabla.column("id", width=45, anchor="center", stretch=False)
        tabla.column("nombre", width=145, anchor="w")
        tabla.column("telefono", width=105, anchor="center")
        tabla.column("correo", width=175, anchor="w")
        tabla.column("direccion", width=190, anchor="w")
        tabla.column("estado", width=85, anchor="center")

        scroll_tabla = ttk.Scrollbar(
            marco_tabla,
            orient="vertical",
            command=tabla.yview
        )

        tabla.configure(yscrollcommand=scroll_tabla.set)

        scroll_tabla.pack(side="right", fill="y")
        tabla.pack(side="left", fill="both", expand=True)

        tabla.tag_configure("inactivo", foreground=VINO)

        def actualizar_resumen():
            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:
                cursor = conexion.cursor()

                cursor.execute("""
                    SELECT
                        COUNT(*) AS total,
                        COALESCE(SUM(CASE
                            WHEN activo = 1 THEN 1
                            ELSE 0
                        END), 0) AS activos,
                        COALESCE(SUM(CASE
                            WHEN activo = 0 THEN 1
                            ELSE 0
                        END), 0) AS inactivos
                    FROM proveedor
                """)

                total, activos, inactivos = cursor.fetchone()

                total_proveedores.config(text=str(total))
                proveedores_activos.config(text=str(activos))
                proveedores_inactivos.config(text=str(inactivos))

            except mysql.connector.Error as error:
                messagebox.showerror(
                    "Error",
                    f"No se pudo actualizar el resumen.\n\n{error}",
                    parent=ventana
                )

            finally:
                if cursor:
                    cursor.close()

                conexion.close()

        def cargar_proveedores():
            for elemento in tabla.get_children():
                tabla.delete(elemento)

            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:
                cursor = conexion.cursor()

                consulta = """
                    SELECT
                        id_proveedor,
                        nombre,
                        telefono,
                        correo,
                        direccion,
                        activo
                    FROM proveedor
                    WHERE 1 = 1
                """

                parametros = []
                busqueda = entrada_busqueda.get().strip()

                if busqueda:
                    consulta += """
                        AND (
                            nombre LIKE %s
                            OR telefono LIKE %s
                            OR correo LIKE %s
                            OR direccion LIKE %s
                        )
                    """

                    texto = f"%{busqueda}%"
                    parametros.extend([texto, texto, texto, texto])

                estado_filtro = filtro_estado.get()

                if estado_filtro == "Activos":
                    consulta += " AND activo = 1"
                elif estado_filtro == "Inactivos":
                    consulta += " AND activo = 0"

                consulta += " ORDER BY nombre ASC"

                cursor.execute(consulta, tuple(parametros))
                proveedores = cursor.fetchall()

                for proveedor in proveedores:
                    (
                        id_proveedor,
                        nombre,
                        telefono,
                        correo,
                        direccion,
                        activo
                    ) = proveedor

                    estado_texto = "Activo" if activo else "Inactivo"

                    valores = (
                        id_proveedor,
                        nombre or "",
                        telefono or "",
                        correo or "",
                        direccion or "",
                        estado_texto
                    )

                    if activo:
                        tabla.insert("", "end", values=valores)
                    else:
                        tabla.insert(
                            "",
                            "end",
                            values=valores,
                            tags=("inactivo",)
                        )

            except mysql.connector.Error as error:
                messagebox.showerror(
                    "Error de proveedores",
                    f"No se pudieron cargar los proveedores.\n\n{error}",
                    parent=ventana
                )

            finally:
                if cursor:
                    cursor.close()

                conexion.close()

            actualizar_resumen()

        def abrir_formulario(proveedor=None):
            formulario = tk.Toplevel(ventana)
            es_edicion = proveedor is not None

            formulario.title(
                "Editar proveedor" if es_edicion
                else "Registrar proveedor"
            )

            formulario.geometry("480x490")
            formulario.configure(bg=CREMA)
            formulario.resizable(False, False)
            formulario.transient(ventana)
            formulario.grab_set()

            try:
                formulario.iconbitmap(LOGO_ICO)
            except Exception:
                pass

            tk.Label(
                formulario,
                text="Editar proveedor" if es_edicion else "Nuevo proveedor",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Georgia", 20, "bold")
            ).pack(pady=(22, 4))

            tk.Label(
                formulario,
                text="Información de contacto",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(pady=(0, 15))

            campos = tk.Frame(formulario, bg=CREMA)
            campos.pack(fill="x", padx=35)

            def crear_campo(etiqueta):
                tk.Label(
                    campos,
                    text=etiqueta,
                    bg=CREMA,
                    fg=VERDE_OSCURO,
                    font=("Segoe UI", 9, "bold")
                ).pack(anchor="w", pady=(5, 3))

                entrada = tk.Entry(
                    campos,
                    bg=BLANCO,
                    fg=VERDE_OSCURO,
                    font=("Segoe UI", 10),
                    relief="flat",
                    highlightbackground=BEIGE,
                    highlightthickness=1
                )

                entrada.pack(fill="x", ipady=6)
                return entrada

            entrada_nombre = crear_campo("Nombre del proveedor")
            entrada_telefono = crear_campo("Teléfono")
            entrada_correo = crear_campo("Correo electrónico")
            entrada_direccion = crear_campo("Dirección")

            if es_edicion:
                entrada_nombre.insert(0, proveedor[1] or "")
                entrada_telefono.insert(0, proveedor[2] or "")
                entrada_correo.insert(0, proveedor[3] or "")
                entrada_direccion.insert(0, proveedor[4] or "")

            def guardar_proveedor():
                nombre = entrada_nombre.get().strip()
                telefono = entrada_telefono.get().strip()
                correo = entrada_correo.get().strip()
                direccion = entrada_direccion.get().strip()

                if not nombre:
                    messagebox.showwarning(
                        "Campo obligatorio",
                        "Escribe el nombre del proveedor.",
                        parent=formulario
                    )
                    entrada_nombre.focus_set()
                    return

                if correo and (
                    "@" not in correo
                    or "." not in correo.split("@")[-1]
                ):
                    messagebox.showwarning(
                        "Correo inválido",
                        "Escribe una dirección de correo válida.",
                        parent=formulario
                    )
                    entrada_correo.focus_set()
                    return

                conexion = conectar()

                if not conexion:
                    return

                cursor = None

                try:
                    cursor = conexion.cursor()

                    if es_edicion:
                        cursor.execute("""
                            UPDATE proveedor
                            SET nombre = %s, telefono = %s,
                                correo = %s, direccion = %s
                            WHERE id_proveedor = %s
                        """, (
                            nombre,
                            telefono or None,
                            correo or None,
                            direccion or None,
                            proveedor[0]
                        ))

                        mensaje = "Proveedor actualizado correctamente."

                    else:
                        cursor.execute("""
                            INSERT INTO proveedor (
                                nombre,
                                telefono,
                                correo,
                                direccion,
                                activo
                            )
                            VALUES (%s, %s, %s, %s, 1)
                        """, (
                            nombre,
                            telefono or None,
                            correo or None,
                            direccion or None
                        ))

                        mensaje = "Proveedor registrado correctamente."

                    conexion.commit()
                    formulario.destroy()

                    cargar_proveedores()

                    messagebox.showinfo(
                        "Operación exitosa",
                        mensaje,
                        parent=ventana
                    )

                except mysql.connector.Error as error:
                    conexion.rollback()

                    messagebox.showerror(
                        "Error al guardar",
                        f"No se pudo guardar el proveedor.\n\n{error}",
                        parent=formulario
                    )

                finally:
                    if cursor:
                        cursor.close()

                    conexion.close()

            marco_botones = tk.Frame(formulario, bg=CREMA)
            marco_botones.pack(pady=20)

            tk.Button(
                marco_botones,
                text="Guardar",
                command=guardar_proveedor,
                bg=VERDE,
                fg=BEIGE,
                activebackground=VERDE_OSCURO,
                activeforeground=BEIGE,
                relief="flat",
                bd=0,
                font=("Segoe UI", 10, "bold"),
                padx=25,
                pady=9,
                cursor="hand2"
            ).pack(side="left", padx=5)

            tk.Button(
                marco_botones,
                text="Cancelar",
                command=formulario.destroy,
                bg=BEIGE,
                fg=VERDE_OSCURO,
                activebackground=BEIGE,
                activeforeground=VERDE_OSCURO,
                relief="flat",
                bd=0,
                font=("Segoe UI", 10, "bold"),
                padx=25,
                pady=9,
                cursor="hand2"
            ).pack(side="left", padx=5)

        def obtener_proveedor_seleccionado():
            seleccion = tabla.selection()

            if not seleccion:
                messagebox.showwarning(
                    "Seleccionar proveedor",
                    "Selecciona un proveedor de la tabla.",
                    parent=ventana
                )
                return None

            valores = tabla.item(seleccion[0], "values")
            return valores

        def editar_proveedor():
            proveedor = obtener_proveedor_seleccionado()

            if proveedor is None:
                return

            if proveedor[5] == "Inactivo":
                messagebox.showwarning(
                    "Proveedor inactivo",
                    "Activa el proveedor antes de editarlo.",
                    parent=ventana
                )
                return

            abrir_formulario(proveedor)

        def eliminar_proveedor():
            proveedor = obtener_proveedor_seleccionado()

            if proveedor is None:
                return

            confirmar = messagebox.askyesno(
                "Confirmar eliminación",
                f"¿Deseas eliminar al proveedor «{proveedor[1]}»?\n\n"
                "Esta acción no se puede deshacer.",
                parent=ventana
            )

            if not confirmar:
                return

            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:
                cursor = conexion.cursor()

                cursor.execute("""
                    DELETE FROM proveedor
                    WHERE id_proveedor = %s
                """, (proveedor[0],))

                conexion.commit()

                messagebox.showinfo(
                    "Proveedor eliminado",
                    "El proveedor se eliminó correctamente.",
                    parent=ventana
                )

                cargar_proveedores()

            except mysql.connector.Error as error:
                conexion.rollback()

                messagebox.showerror(
                    "No se puede eliminar",
                    "No fue posible eliminar el proveedor.\n\n"
                    "Es posible que tenga ingredientes relacionados "
                    "en el inventario.\n\n"
                    f"Detalle: {error}",
                    parent=ventana
                )

            finally:
                if cursor:
                    cursor.close()

                conexion.close()

        def cambiar_estado_proveedor():
            proveedor = obtener_proveedor_seleccionado()

            if proveedor is None:
                return

            id_proveedor = proveedor[0]
            nombre = proveedor[1]
            esta_activo = proveedor[5] == "Activo"

            nuevo_estado = 0 if esta_activo else 1
            accion = "desactivar" if esta_activo else "activar"

            confirmar = messagebox.askyesno(
                "Confirmar operación",
                f"¿Deseas {accion} al proveedor «{nombre}»?",
                parent=ventana
            )

            if not confirmar:
                return

            conexion = conectar()

            if not conexion:
                return

            cursor = None

            try:
                cursor = conexion.cursor()

                cursor.execute("""
                    UPDATE proveedor
                    SET activo = %s
                    WHERE id_proveedor = %s
                """, (nuevo_estado, id_proveedor))

                conexion.commit()

                mensaje = (
                    "Proveedor desactivado correctamente."
                    if nuevo_estado == 0
                    else "Proveedor activado correctamente."
                )

                messagebox.showinfo(
                    "Operación exitosa",
                    mensaje,
                    parent=ventana
                )

                cargar_proveedores()

            except mysql.connector.Error as error:
                conexion.rollback()

                messagebox.showerror(
                    "Error",
                    f"No se pudo cambiar el estado del proveedor.\n\n{error}",
                    parent=ventana
                )

            finally:
                if cursor:
                    cursor.close()

                conexion.close()

        tk.Button(
            barra_acciones,
            text="+ Nuevo proveedor",
            command=lambda: abrir_formulario(),
            bg=VERDE,
            fg=BEIGE,
            activebackground=VERDE_OSCURO,
            activeforeground=BEIGE,
            relief="flat",
            bd=0,
            font=("Segoe UI", 9, "bold"),
            padx=15,
            pady=10,
            cursor="hand2"
        ).pack(side="left")

        tk.Button(
            barra_acciones,
            text="Editar",
            command=editar_proveedor,
            bg=BEIGE,
            fg=VERDE_OSCURO,
            activebackground=BEIGE,
            activeforeground=VERDE_OSCURO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 9, "bold"),
            padx=20,
            pady=10,
            cursor="hand2"
        ).pack(side="left", padx=8)

        tk.Button(
            barra_acciones,
            text="Eliminar",
            command=eliminar_proveedor,
            bg=VINO,
            fg=BLANCO,
            activebackground=VINO,
            activeforeground=BLANCO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 9, "bold"),
            padx=15,
            pady=10,
            cursor="hand2"
        ).pack(side="left")

        entrada_busqueda.bind(
            "<KeyRelease>",
            lambda evento: cargar_proveedores()
        )

        filtro_estado.bind(
            "<<ComboboxSelected>>",
            lambda evento: cargar_proveedores()
        )

        actualizar_resumen()
        cargar_proveedores()


    def vista_pendiente(nombre):

        limpiar_vista()
        activar_boton(nombre)

        tk.Label(
            vista_actual,
            text=nombre.title(),
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Georgia", 22, "bold")
        ).pack(
            anchor="w",
            pady=(20, 5)
        )

        tk.Label(
            vista_actual,
            text="Esta sección se desarrollará a continuación.",
            bg=CREMA,
            fg=GRIS,
            font=("Segoe UI", 10)
        ).pack(
            anchor="w"
        )


    def mostrar_menu():
        from tkinter import filedialog
        from PIL import Image, ImageTk

        limpiar_vista()
        mostrar_encabezados(False)
        activar_boton("MENÚ")

        CARPETA_BASE = os.path.dirname(os.path.abspath(__file__))
        CARPETA_IMAGENES = os.path.join(CARPETA_BASE, "imagenes_menu")
        os.makedirs(CARPETA_IMAGENES, exist_ok=True)

        categorias_base = [
            "Antipasti",
            "Primi",
            "Secondi",
            "Dolci & Vini"
        ]

        imagenes_cache = []
        platillos = []

        encabezado = tk.Frame(vista_actual, bg=CREMA)
        encabezado.pack(fill="x", pady=(8, 5))

        tk.Label(
            encabezado,
            text="Menú del Restaurante",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Georgia", 23, "bold")
        ).pack(anchor="w")

        tk.Label(
            encabezado,
            text="Gestión de platillos, precios y disponibilidad",
            bg=CREMA,
            fg=GRIS,
            font=("Segoe UI", 9)
        ).pack(side="left", pady=(8, 0))

        barra_busqueda = tk.Frame(vista_actual, bg=CREMA)
        barra_busqueda.pack(fill="x", pady=(12, 12))

        tk.Label(
            barra_busqueda,
            text="Buscar:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left")

        busqueda = tk.StringVar()

        entrada_busqueda = tk.Entry(
            barra_busqueda,
            textvariable=busqueda,
            font=("Segoe UI", 10),
            bg=BLANCO,
            fg=VERDE_OSCURO,
            relief="flat",
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        entrada_busqueda.pack(
            side="left",
            padx=(8, 12),
            ipady=6,
            fill="x",
            expand=True
        )

        tk.Label(
            barra_busqueda,
            text="Disponibilidad:",
            bg=CREMA,
            fg=VERDE_OSCURO,
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 6))

        filtro_disponibilidad = ttk.Combobox(
            barra_busqueda,
            values=["Todos", "Disponible", "Agotado"],
            state="readonly",
            width=13,
            font=("Segoe UI", 9)
        )
        filtro_disponibilidad.set("Todos")
        filtro_disponibilidad.pack(side="left")

        barra_acciones = tk.Frame(vista_actual, bg=CREMA)
        barra_acciones.pack(fill="x", pady=(0, 12))

        tk.Button(
            barra_acciones,
            text="+ Nuevo platillo",
            command=lambda: abrir_formulario(),
            bg=VERDE,
            fg=BEIGE,
            activebackground=VERDE_OSCURO,
            activeforeground=BEIGE,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2",
            padx=16,
            pady=9
        ).pack(side="left")

        tk.Button(
            barra_acciones,
            text="Actualizar",
            command=lambda: cargar_platillos(),
            bg=BEIGE,
            fg=VERDE_OSCURO,
            activebackground=BEIGE,
            activeforeground=VERDE_OSCURO,
            relief="flat",
            bd=0,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2",
            padx=16,
            pady=9
        ).pack(side="left", padx=8)

        marco_principal = tk.Frame(
            vista_actual,
            bg=CREMA_CLARO,
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        marco_principal.pack(fill="both", expand=True)

        canvas = tk.Canvas(
            marco_principal,
            bg=CREMA_CLARO,
            highlightthickness=0
        )

        scrollbar = ttk.Scrollbar(
            marco_principal,
            orient="vertical",
            command=canvas.yview
        )

        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        contenedor_tarjetas = tk.Frame(canvas, bg=CREMA_CLARO)

        ventana_tarjetas = canvas.create_window(
            (0, 0),
            window=contenedor_tarjetas,
            anchor="nw"
        )

        def actualizar_scroll(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def ajustar_ancho(event):
            canvas.itemconfigure(
                ventana_tarjetas,
                width=event.width
            )

        contenedor_tarjetas.bind("<Configure>", actualizar_scroll)
        canvas.bind("<Configure>", ajustar_ancho)

        def desplazar_tarjetas(event):
            if event.delta:
                canvas.yview_scroll(
                    int(-event.delta / 120),
                    "units"
                )

        def activar_scroll_recursivo(widget):
            widget.bind("<MouseWheel>", desplazar_tarjetas)

            for hijo in widget.winfo_children():
                activar_scroll_recursivo(hijo)

        def cargar_platillos():
            nonlocal platillos

            conexion = None
            cursor = None

            try:
                conexion = conectar()

                if conexion is None:
                    return

                cursor = conexion.cursor(dictionary=True)

                cursor.execute("""
                    SELECT
                        id_platillo,
                        nombre,
                        descripcion,
                        precio,
                        categoria,
                        disponibilidad,
                        popularidad,
                        resena_promedio,
                        imagen_url,
                        vegetariano,
                        sin_gluten,
                        distintivo
                    FROM Platillo
                    ORDER BY
                        FIELD(
                            categoria,
                            'Antipasti',
                            'Primi',
                            'Secondi',
                            'Dolci & Vini'
                        ),
                        nombre ASC
                """)

                platillos = cursor.fetchall()

            except Exception as error:
                messagebox.showerror(
                    "Error al cargar el menú",
                    f"No se pudieron cargar los platillos.\n\n{error}"
                )
                return

            finally:
                if cursor is not None:
                    cursor.close()

                if conexion is not None:
                    conexion.close()

            renderizar_menu()

        def abrir_formulario(platillo=None):
            es_edicion = platillo is not None

            formulario = tk.Toplevel(vista_actual.winfo_toplevel())
            formulario.title(
                "Editar platillo" if es_edicion else "Nuevo platillo"
            )
            formulario.geometry("560x720")
            formulario.minsize(460, 500)
            formulario.configure(bg=CREMA)
            formulario.transient(vista_actual.winfo_toplevel())
            formulario.grab_set()

            ruta_icono = os.path.join(CARPETA_BASE, "vivaldi.ico")
            if os.path.isfile(ruta_icono):
                try:
                    formulario.iconbitmap(ruta_icono)
                except tk.TclError:
                    pass

            tk.Label(
                formulario,
                text="Editar platillo" if es_edicion else "Nuevo platillo",
                bg=CREMA,
                fg=VERDE_OSCURO,
                font=("Georgia", 21, "bold")
            ).pack(anchor="w", padx=22, pady=(18, 4))

            tk.Label(
                formulario,
                text="Completa la información del platillo",
                bg=CREMA,
                fg=GRIS,
                font=("Segoe UI", 9)
            ).pack(anchor="w", padx=22, pady=(0, 12))

            area = tk.Frame(formulario, bg=CREMA)
            area.pack(fill="both", expand=True, padx=22, pady=(0, 10))

            canvas_form = tk.Canvas(
                area,
                bg=CREMA,
                highlightthickness=0
            )
            scroll_form = ttk.Scrollbar(
                area,
                orient="vertical",
                command=canvas_form.yview
            )

            canvas_form.configure(yscrollcommand=scroll_form.set)
            scroll_form.pack(side="right", fill="y")
            canvas_form.pack(side="left", fill="both", expand=True)

            contenido_form = tk.Frame(canvas_form, bg=CREMA_CLARO)
            ventana_form = canvas_form.create_window(
                (0, 0),
                window=contenido_form,
                anchor="nw"
            )

            contenido_form.bind(
                "<Configure>",
                lambda e: canvas_form.configure(
                    scrollregion=canvas_form.bbox("all")
                )
            )

            canvas_form.bind(
                "<Configure>",
                lambda e: canvas_form.itemconfigure(
                    ventana_form,
                    width=e.width
                )
            )

            def crear_campo(etiqueta, valor=""):
                tk.Label(
                    contenido_form,
                    text=etiqueta,
                    bg=CREMA_CLARO,
                    fg=VERDE_OSCURO,
                    font=("Segoe UI", 9, "bold")
                ).pack(anchor="w", padx=16, pady=(9, 4))

                entrada = tk.Entry(
                    contenido_form,
                    font=("Segoe UI", 10),
                    bg=BLANCO,
                    fg=VERDE_OSCURO,
                    relief="flat",
                    highlightbackground=BEIGE,
                    highlightthickness=1
                )
                entrada.pack(fill="x", padx=16, ipady=6)
                entrada.insert(0, str(valor or ""))
                return entrada

            nombre_var = crear_campo(
                "Nombre del platillo",
                platillo.get("nombre") if es_edicion else ""
            )

            tk.Label(
                contenido_form,
                text="Descripción",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(9, 4))

            descripcion_var = tk.Text(
                contenido_form,
                height=3,
                font=("Segoe UI", 10),
                bg=BLANCO,
                fg=VERDE_OSCURO,
                relief="flat",
                wrap="word"
            )
            descripcion_var.pack(fill="x", padx=16)

            if es_edicion and platillo.get("descripcion"):
                descripcion_var.insert(
                    "1.0",
                    platillo["descripcion"]
                )

            precio_var = crear_campo(
                "Precio ($)",
                platillo.get("precio") if es_edicion else ""
            )

            tk.Label(
                contenido_form,
                text="Categoría",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(9, 4))

            categoria_var = tk.StringVar(
                value=(
                    platillo.get("categoria")
                    if es_edicion
                    else categorias_base[0]
                )
            )

            categorias_formulario = list(categorias_base)

            if (
                es_edicion
                and platillo.get("categoria")
                and platillo["categoria"] not in categorias_formulario
            ):
                categorias_formulario.append(platillo["categoria"])

            ttk.Combobox(
                contenido_form,
                textvariable=categoria_var,
                values=categorias_formulario,
                state="readonly"
            ).pack(fill="x", padx=16, ipady=4)

            tk.Label(
                contenido_form,
                text="Disponibilidad",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(9, 4))

            disponibilidad_var = tk.StringVar(
                value=(
                    platillo.get("disponibilidad")
                    if es_edicion
                    else "Disponible"
                ) or "Disponible"
            )

            ttk.Combobox(
                contenido_form,
                textvariable=disponibilidad_var,
                values=["Disponible", "Agotado"],
                state="readonly"
            ).pack(fill="x", padx=16, ipady=4)

            vegetariano_var = tk.BooleanVar(
                value=bool(platillo.get("vegetariano"))
                if es_edicion else False
            )

            sin_gluten_var = tk.BooleanVar(
                value=bool(platillo.get("sin_gluten"))
                if es_edicion else False
            )

            tk.Label(
                contenido_form,
                text="Características alimentarias",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(9, 4))

            tk.Checkbutton(
                contenido_form,
                text="Vegetariano",
                variable=vegetariano_var,
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                activebackground=CREMA_CLARO,
                selectcolor=BLANCO
            ).pack(anchor="w", padx=16)

            tk.Checkbutton(
                contenido_form,
                text="Sin gluten",
                variable=sin_gluten_var,
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                activebackground=CREMA_CLARO,
                selectcolor=BLANCO
            ).pack(anchor="w", padx=16)

            tk.Label(
                contenido_form,
                text="Distintivo",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(9, 4))

            distintivo_var = tk.StringVar(
                value=(
                    platillo.get("distintivo") or ""
                    if es_edicion else ""
                )
            )

            ttk.Combobox(
                contenido_form,
                textvariable=distintivo_var,
                values=["", "Favorito", "Destacado"],
                state="readonly"
            ).pack(fill="x", padx=16, ipady=4)

            ruta_imagen = {
                "ruta": platillo.get("imagen_url") or ""
                if es_edicion else ""
            }

            tk.Label(
                contenido_form,
                text="Fotografía del platillo",
                bg=CREMA_CLARO,
                fg=VERDE_OSCURO,
                font=("Segoe UI", 9, "bold")
            ).pack(anchor="w", padx=16, pady=(12, 4))

            etiqueta_foto = tk.Label(
                contenido_form,
                text="Sin fotografía seleccionada",
                bg=CREMA_CLARO,
                fg=GRIS
            )
            etiqueta_foto.pack(anchor="w", padx=16)

            foto_form_cache = {"imagen": None}

            def mostrar_foto_form(ruta):
                ruta_completa = ruta

                if ruta and not os.path.isabs(ruta):
                    ruta_completa = os.path.join(CARPETA_BASE, ruta)

                if ruta_completa and os.path.isfile(ruta_completa):
                    try:
                        with Image.open(ruta_completa) as imagen:
                            imagen = imagen.convert("RGB")
                            imagen.thumbnail((220, 110))
                            foto = ImageTk.PhotoImage(imagen)

                        foto_form_cache["imagen"] = foto
                        etiqueta_foto.configure(image=foto, text="")
                    except Exception:
                        etiqueta_foto.configure(
                            image="",
                            text="No se pudo mostrar la fotografía"
                        )
                else:
                    etiqueta_foto.configure(
                        image="",
                        text="Sin fotografía seleccionada"
                    )

            if ruta_imagen["ruta"]:
                mostrar_foto_form(ruta_imagen["ruta"])

            def seleccionar_fotografia():
                archivo = filedialog.askopenfilename(
                    parent=formulario,
                    title="Seleccionar fotografía",
                    filetypes=[
                        ("Imágenes", "*.png *.jpg *.jpeg *.gif *.bmp *.webp"),
                        ("Todos los archivos", "*.*")
                    ]
                )

                if not archivo:
                    return

                try:
                    with Image.open(archivo) as imagen:
                        imagen.verify()

                    ruta_imagen["ruta"] = archivo
                    mostrar_foto_form(archivo)

                except Exception as error:
                    messagebox.showerror(
                        "Error de imagen",
                        f"No se pudo abrir la fotografía.\n\n{error}",
                        parent=formulario
                    )

            tk.Button(
                contenido_form,
                text="Seleccionar fotografía…",
                command=seleccionar_fotografia,
                bg=BEIGE,
                fg=VERDE_OSCURO,
                activebackground=BEIGE,
                relief="flat",
                cursor="hand2",
                padx=12,
                pady=7
            ).pack(anchor="w", padx=16, pady=(6, 14))

            def guardar_platillo():
                nombre = nombre_var.get().strip()
                descripcion = descripcion_var.get("1.0", "end").strip()
                categoria = categoria_var.get().strip()
                disponibilidad = disponibilidad_var.get().strip()
                distintivo = distintivo_var.get().strip() or None

                if not nombre or not precio_var.get().strip() or not categoria:
                    messagebox.showwarning(
                        "Campos obligatorios",
                        "Ingresa el nombre, el precio y la categoría.",
                        parent=formulario
                    )
                    return

                try:
                    precio = float(
                        precio_var.get().strip().replace(",", ".")
                    )
                    if precio < 0:
                        raise ValueError
                except ValueError:
                    messagebox.showwarning(
                        "Precio inválido",
                        "Introduce un precio numérico válido.",
                        parent=formulario
                    )
                    return

                imagen_final = (
                    platillo.get("imagen_url")
                    if es_edicion else None
                )

                imagen_seleccionada = ruta_imagen["ruta"]

                if imagen_seleccionada:
                    es_imagen_local = (
                        os.path.isabs(imagen_seleccionada)
                        and os.path.dirname(imagen_seleccionada)
                        == os.path.abspath(CARPETA_IMAGENES)
                    )

                    if es_imagen_local:
                        imagen_final = os.path.join(
                            "imagenes_menu",
                            os.path.basename(imagen_seleccionada)
                        )
                    elif (
                        es_edicion
                        and imagen_seleccionada
                        == platillo.get("imagen_url")
                    ):
                        imagen_final = platillo.get("imagen_url")
                    else:
                        extension = os.path.splitext(
                            imagen_seleccionada
                        )[1].lower()

                        if extension not in [
                            ".png", ".jpg", ".jpeg",
                            ".gif", ".bmp", ".webp"
                        ]:
                            messagebox.showwarning(
                                "Formato no válido",
                                "Selecciona una imagen PNG, JPG, JPEG, GIF, BMP o WEBP.",
                                parent=formulario
                            )
                            return

                        nombre_archivo = f"{uuid.uuid4().hex}{extension}"
                        destino = os.path.join(
                            CARPETA_IMAGENES,
                            nombre_archivo
                        )

                        try:
                            shutil.copy2(imagen_seleccionada, destino)
                            imagen_final = os.path.join(
                                "imagenes_menu",
                                nombre_archivo
                            )
                        except Exception as error:
                            messagebox.showerror(
                                "Error al guardar la fotografía",
                                str(error),
                                parent=formulario
                            )
                            return

                conexion = None
                cursor = None

                try:
                    conexion = conectar()

                    if conexion is None:
                        return

                    cursor = conexion.cursor()

                    if es_edicion:
                        cursor.execute("""
                            UPDATE Platillo
                            SET nombre = %s,
                                descripcion = %s,
                                precio = %s,
                                categoria = %s,
                                disponibilidad = %s,
                                imagen_url = %s,
                                vegetariano = %s,
                                sin_gluten = %s,
                                distintivo = %s
                            WHERE id_platillo = %s
                        """, (
                            nombre,
                            descripcion,
                            precio,
                            categoria,
                            disponibilidad,
                            imagen_final,
                            vegetariano_var.get(),
                            sin_gluten_var.get(),
                            distintivo,
                            platillo["id_platillo"]
                        ))
                    else:
                        cursor.execute("""
                            INSERT INTO Platillo (
                                nombre,
                                descripcion,
                                precio,
                                categoria,
                                disponibilidad,
                                popularidad,
                                resena_promedio,
                                imagen_url,
                                vegetariano,
                                sin_gluten,
                                distintivo
                            )
                            VALUES (
                                %s, %s, %s, %s, %s,
                                %s, %s, %s, %s, %s, %s
                            )
                        """, (
                            nombre,
                            descripcion,
                            precio,
                            categoria,
                            disponibilidad,
                            0,
                            0.0,
                            imagen_final,
                            vegetariano_var.get(),
                            sin_gluten_var.get(),
                            distintivo
                        ))

                    conexion.commit()
                    formulario.destroy()
                    cargar_platillos()

                    messagebox.showinfo(
                        "Menú actualizado",
                        "El platillo se guardó correctamente."
                    )

                except Exception as error:
                    if conexion is not None:
                        conexion.rollback()

                    messagebox.showerror(
                        "Error al guardar",
                        f"No se pudo guardar el platillo.\n\n{error}",
                        parent=formulario
                    )

                finally:
                    if cursor is not None:
                        cursor.close()

                    if conexion is not None:
                        conexion.close()

            botones_form = tk.Frame(formulario, bg=CREMA)
            botones_form.pack(fill="x", padx=22, pady=(0, 16))

            tk.Button(
                botones_form,
                text="Guardar cambios" if es_edicion else "Guardar platillo",
                command=guardar_platillo,
                bg=VERDE,
                fg=BEIGE,
                activebackground=VERDE_OSCURO,
                activeforeground=BEIGE,
                relief="flat",
                font=("Segoe UI", 10, "bold"),
                cursor="hand2",
                padx=15,
                pady=9
            ).pack(side="left")

            tk.Button(
                botones_form,
                text="Cancelar",
                command=formulario.destroy,
                bg=VINO,
                fg=BLANCO,
                activebackground=VINO,
                activeforeground=BLANCO,
                relief="flat",
                font=("Segoe UI", 10, "bold"),
                cursor="hand2",
                padx=15,
                pady=9
            ).pack(side="left", padx=8)

        def eliminar_platillo(platillo):
            confirmar = messagebox.askyesno(
                "Confirmar eliminación",
                f"¿Deseas eliminar el platillo '{platillo['nombre']}'?"
            )

            if not confirmar:
                return

            conexion = None
            cursor = None

            try:
                conexion = conectar()

                if conexion is None:
                    return

                cursor = conexion.cursor()
                cursor.execute(
                    "DELETE FROM Platillo WHERE id_platillo = %s",
                    (platillo["id_platillo"],)
                )
                conexion.commit()

                cargar_platillos()

                messagebox.showinfo(
                    "Platillo eliminado",
                    "El platillo fue eliminado correctamente."
                )

            except Exception as error:
                if conexion is not None:
                    conexion.rollback()

                messagebox.showerror(
                    "No se pudo eliminar",
                    "El platillo podría estar relacionado con pedidos "
                    "u otros registros.\n\n"
                    f"Detalle: {error}"
                )

            finally:
                if cursor is not None:
                    cursor.close()

                if conexion is not None:
                    conexion.close()

        def renderizar_menu(*args):
            for widget in contenedor_tarjetas.winfo_children():
                widget.destroy()

            imagenes_cache.clear()

            texto = busqueda.get().strip().lower()
            disponibilidad_filtro = filtro_disponibilidad.get()

            resultados = []

            for p in platillos:
                nombre = str(p.get("nombre") or "")
                categoria = str(p.get("categoria") or "")
                descripcion = str(p.get("descripcion") or "")
                disponibilidad = str(p.get("disponibilidad") or "")

                coincide_texto = (
                        texto in nombre.lower()
                        or texto in categoria.lower()
                        or texto in descripcion.lower()
                )

                coincide_estado = (
                        disponibilidad_filtro == "Todos"
                        or disponibilidad.lower()
                        == disponibilidad_filtro.lower()
                )

                if coincide_texto and coincide_estado:
                    resultados.append(p)

            if not resultados:
                tk.Label(
                    contenedor_tarjetas,
                    text="No se encontraron platillos.",
                    bg=CREMA_CLARO,
                    fg=GRIS,
                    font=("Segoe UI", 12),
                    pady=35
                ).grid(
                    row=0,
                    column=0,
                    columnspan=2,
                    sticky="ew",
                    padx=20,
                    pady=25
                )

                canvas.after_idle(actualizar_scroll)
                return

            for columna in range(2):
                contenedor_tarjetas.grid_columnconfigure(
                    columna,
                    weight=1,
                    uniform="columna_platillos"
                )

            categorias_menu = [
                "Antipasti",
                "Primi",
                "Secondi",
                "Dolci & Vini"
            ]

            agrupados = {
                categoria: [] for categoria in categorias_menu
            }

            otras_categorias = []

            for p in resultados:
                categoria_platillo = str(
                    p.get("categoria") or ""
                ).strip()

                categoria_encontrada = None

                for categoria in categorias_menu:
                    if (
                            categoria_platillo.casefold()
                            == categoria.casefold()
                    ):
                        categoria_encontrada = categoria
                        break

                if categoria_encontrada:
                    agrupados[categoria_encontrada].append(p)
                else:
                    otras_categorias.append(p)

            fila_actual = 0

            for categoria in categorias_menu:
                lista = agrupados[categoria]

                if not lista:
                    continue

                encabezado = tk.Frame(
                    contenedor_tarjetas,
                    bg=CREMA_CLARO
                )

                encabezado.grid(
                    row=fila_actual,
                    column=0,
                    columnspan=2,
                    sticky="ew",
                    padx=10,
                    pady=(20, 8)
                )

                tk.Label(
                    encabezado,
                    text=categoria,
                    bg=CREMA_CLARO,
                    fg=VERDE_OSCURO,
                    font=("Georgia", 18, "bold")
                ).pack(side="left")

                tk.Frame(
                    encabezado,
                    bg=BEIGE,
                    height=2
                ).pack(
                    side="left",
                    fill="x",
                    expand=True,
                    padx=(12, 0),
                    pady=7
                )

                fila_actual += 1

                contenedor_categoria = tk.Frame(
                    contenedor_tarjetas,
                    bg=CREMA_CLARO
                )

                contenedor_categoria.grid(
                    row=fila_actual,
                    column=0,
                    columnspan=2,
                    sticky="ew"
                )

                contenedor_categoria.grid_columnconfigure(
                    0,
                    weight=1,
                    uniform="tarjetas_categoria"
                )

                contenedor_categoria.grid_columnconfigure(
                    1,
                    weight=1,
                    uniform="tarjetas_categoria"
                )

                for indice, platillo in enumerate(lista):
                    fila = indice // 2
                    columna = indice % 2

                    tarjeta = tk.Frame(
                        contenedor_categoria,
                        bg=BLANCO,
                        highlightbackground=GRIS_CLARO,
                        highlightthickness=1
                    )

                    tarjeta.grid(
                        row=fila,
                        column=columna,
                        sticky="nsew",
                        padx=5,
                        pady=5
                    )

                    tk.Frame(
                        tarjeta,
                        bg=VERDE,
                        height=4
                    ).pack(fill="x")

                    contenido = tk.Frame(
                        tarjeta,
                        bg=BLANCO
                    )

                    contenido.pack(
                        fill="both",
                        expand=True,
                        padx=9,
                        pady=9
                    )

                    cuerpo = tk.Frame(
                        contenido,
                        bg=BLANCO
                    )

                    cuerpo.pack(
                        fill="both",
                        expand=True
                    )

                    marco_imagen = tk.Frame(
                        cuerpo,
                        bg=CREMA_CLARO,
                        width=100,
                        height=110
                    )

                    marco_imagen.pack(
                        side="left",
                        padx=(0, 9)
                    )

                    marco_imagen.pack_propagate(False)

                    ruta = str(
                        platillo.get("imagen_url") or ""
                    ).strip()

                    ruta_completa = ruta

                    if ruta and not os.path.isabs(ruta):
                        ruta_completa = os.path.join(
                            CARPETA_BASE,
                            ruta
                        )

                    foto_cargada = False

                    if (
                            ruta_completa
                            and os.path.isfile(ruta_completa)
                            and Image is not None
                            and ImageTk is not None
                    ):
                        try:
                            with Image.open(ruta_completa) as imagen:
                                imagen = imagen.convert("RGB")
                                imagen.thumbnail((100, 110))
                                foto = ImageTk.PhotoImage(imagen)

                            imagenes_cache.append(foto)

                            tk.Label(
                                marco_imagen,
                                image=foto,
                                bg=CREMA_CLARO
                            ).pack(expand=True)

                            foto_cargada = True

                        except Exception:
                            foto_cargada = False

                    if not foto_cargada:
                        tk.Label(
                            marco_imagen,
                            text="Sin imagen",
                            bg=CREMA_CLARO,
                            fg=GRIS,
                            font=("Segoe UI", 8),
                            justify="center"
                        ).pack(
                            expand=True,
                            fill="both"
                        )

                    info = tk.Frame(
                        cuerpo,
                        bg=BLANCO
                    )

                    info.pack(
                        side="left",
                        fill="both",
                        expand=True
                    )

                    tk.Label(
                        info,
                        text=str(
                            platillo.get("nombre") or "Sin nombre"
                        ),
                        bg=BLANCO,
                        fg=VERDE_OSCURO,
                        font=("Segoe UI", 10, "bold"),
                        wraplength=175,
                        justify="left",
                        anchor="w"
                    ).pack(fill="x")

                    tk.Label(
                        info,
                        text=str(
                            platillo.get("descripcion")
                            or "Sin descripción registrada."
                        ),
                        bg=BLANCO,
                        fg=GRIS,
                        font=("Segoe UI", 8),
                        wraplength=175,
                        justify="left",
                        anchor="w"
                    ).pack(
                        fill="x",
                        pady=(4, 5)
                    )

                    etiquetas = []

                    if platillo.get("vegetariano"):
                        etiquetas.append("Vegetariano")

                    if platillo.get("sin_gluten"):
                        etiquetas.append("Sin gluten")

                    if platillo.get("distintivo"):
                        etiquetas.append(
                            str(platillo["distintivo"])
                        )

                    if etiquetas:
                        tk.Label(
                            info,
                            text=" · ".join(etiquetas),
                            bg=BLANCO,
                            fg=VINO,
                            font=("Segoe UI", 8, "bold"),
                            wraplength=175,
                            justify="left",
                            anchor="w"
                        ).pack(
                            fill="x",
                            pady=(0, 4)
                        )

                    try:
                        texto_precio = (
                            f"${float(platillo['precio']):.2f}"
                        )
                    except (TypeError, ValueError):
                        texto_precio = "$0.00"

                    tk.Label(
                        info,
                        text=texto_precio,
                        bg=BLANCO,
                        fg=VINO,
                        font=("Georgia", 13, "bold"),
                        anchor="w"
                    ).pack(
                        fill="x",
                        pady=(2, 4)
                    )

                    disponible = (
                            str(
                                platillo.get("disponibilidad") or ""
                            ).lower() == "disponible"
                    )

                    tk.Label(
                        info,
                        text=(
                            f"● {platillo.get('disponibilidad') or 'Sin estado'}"
                        ),
                        bg=BLANCO,
                        fg=VERDE if disponible else VINO,
                        font=("Segoe UI", 8, "bold"),
                        anchor="w"
                    ).pack(fill="x")

                    botones = tk.Frame(
                        contenido,
                        bg=BLANCO
                    )

                    botones.pack(
                        fill="x",
                        pady=(10, 0)
                    )

                    tk.Button(
                        botones,
                        text="EDITAR",
                        command=lambda p=platillo: abrir_formulario(p),
                        bg=BEIGE,
                        fg=VERDE_OSCURO,
                        activebackground=BEIGE,
                        activeforeground=VERDE_OSCURO,
                        relief="flat",
                        bd=0,
                        font=("Segoe UI", 8, "bold"),
                        cursor="hand2",
                        padx=12,
                        pady=6
                    ).pack(
                        side="left",
                        fill="x",
                        expand=True
                    )

                    tk.Button(
                        botones,
                        text="ELIMINAR",
                        command=lambda p=platillo: eliminar_platillo(p),
                        bg=VINO,
                        fg=BLANCO,
                        activebackground=VINO,
                        activeforeground=BLANCO,
                        relief="flat",
                        bd=0,
                        font=("Segoe UI", 8, "bold"),
                        cursor="hand2",
                        padx=12,
                        pady=6
                    ).pack(
                        side="left",
                        fill="x",
                        expand=True,
                        padx=(6, 0)
                    )

                    activar_scroll_recursivo(tarjeta)

                fila_actual += 1

            if otras_categorias:
                encabezado = tk.Frame(
                    contenedor_tarjetas,
                    bg=CREMA_CLARO
                )

                encabezado.grid(
                    row=fila_actual,
                    column=0,
                    columnspan=2,
                    sticky="ew",
                    padx=10,
                    pady=(20, 8)
                )

                tk.Label(
                    encabezado,
                    text="Otras categorías",
                    bg=CREMA_CLARO,
                    fg=VERDE_OSCURO,
                    font=("Georgia", 18, "bold")
                ).pack(side="left")

                fila_actual += 1

                contenedor_otros = tk.Frame(
                    contenedor_tarjetas,
                    bg=CREMA_CLARO
                )

                contenedor_otros.grid(
                    row=fila_actual,
                    column=0,
                    columnspan=2,
                    sticky="ew"
                )

                contenedor_otros.grid_columnconfigure(
                    0,
                    weight=1,
                    uniform="otras_categorias"
                )

                contenedor_otros.grid_columnconfigure(
                    1,
                    weight=1,
                    uniform="otras_categorias"
                )

                for indice, platillo in enumerate(otras_categorias):
                    fila = indice // 2
                    columna = indice % 2

                    tarjeta = tk.Frame(
                        contenedor_otros,
                        bg=BLANCO,
                        highlightbackground=GRIS_CLARO,
                        highlightthickness=1,
                        padx=10,
                        pady=10
                    )

                    tarjeta.grid(
                        row=fila,
                        column=columna,
                        sticky="nsew",
                        padx=5,
                        pady=5
                    )

                    tk.Label(
                        tarjeta,
                        text=str(
                            platillo.get("nombre") or "Sin nombre"
                        ),
                        bg=BLANCO,
                        fg=VERDE_OSCURO,
                        font=("Segoe UI", 10, "bold"),
                        wraplength=200,
                        justify="left"
                    ).pack(
                        anchor="w",
                        fill="x"
                    )

                    tk.Label(
                        tarjeta,
                        text=str(
                            platillo.get("categoria") or "Sin categoría"
                        ),
                        bg=BLANCO,
                        fg=VINO,
                        font=("Segoe UI", 8, "bold")
                    ).pack(anchor="w", pady=(3, 5))

                    try:
                        texto_precio = (
                            f"${float(platillo['precio']):.2f}"
                        )
                    except (TypeError, ValueError):
                        texto_precio = "$0.00"

                    tk.Label(
                        tarjeta,
                        text=texto_precio,
                        bg=BLANCO,
                        fg=VINO,
                        font=("Georgia", 13, "bold")
                    ).pack(anchor="w", pady=(0, 8))

                    botones = tk.Frame(
                        tarjeta,
                        bg=BLANCO
                    )

                    botones.pack(fill="x")

                    tk.Button(
                        botones,
                        text="EDITAR",
                        command=lambda p=platillo: abrir_formulario(p),
                        bg=BEIGE,
                        fg=VERDE_OSCURO,
                        relief="flat",
                        cursor="hand2",
                        padx=10,
                        pady=5
                    ).pack(side="left", expand=True, fill="x")

                    tk.Button(
                        botones,
                        text="ELIMINAR",
                        command=lambda p=platillo: eliminar_platillo(p),
                        bg=VINO,
                        fg=BLANCO,
                        relief="flat",
                        cursor="hand2",
                        padx=10,
                        pady=5
                    ).pack(
                        side="left",
                        expand=True,
                        fill="x",
                        padx=(6, 0)
                    )

                    activar_scroll_recursivo(tarjeta)

            canvas.after_idle(actualizar_scroll)

        busqueda.trace_add("write", renderizar_menu)

        filtro_disponibilidad.bind(
            "<<ComboboxSelected>>",
            renderizar_menu
        )

        cargar_platillos()

    def obtener_datos_reportes(periodo="Hoy"):
        conexion = conectar()

        if not conexion:
            return {
                "ingresos": 0,
                "ventas": 0,
                "facturas": []
            }

        cursor = None

        try:
            cursor = conexion.cursor()

            filtros = {
                "Hoy": "DATE(f.fecha) = CURDATE()",
                "Esta semana": (
                    "YEARWEEK(f.fecha, 1) = YEARWEEK(CURDATE(), 1)"
                ),
                "Este mes": (
                    "YEAR(f.fecha) = YEAR(CURDATE()) "
                    "AND MONTH(f.fecha) = MONTH(CURDATE())"
                ),
                "Este año": "YEAR(f.fecha) = YEAR(CURDATE())"
            }

            condicion = filtros.get(periodo, filtros["Hoy"])

            cursor.execute(
                f"""
                SELECT
                    f.id_factura,
                    f.fecha,
                    f.total,
                    f.metodo_pago,
                    CONCAT(e.nombre, ' ', e.apellido) AS nombre_mesero
                FROM Factura AS f
                LEFT JOIN Orden AS o
                    ON f.id_orden = o.id_orden
                LEFT JOIN Empleado AS e
                    ON o.id_empleado = e.id_empleado
                WHERE f.estado_pago = 'Pagado'
                  AND {condicion}
                ORDER BY f.fecha DESC, f.id_factura DESC
                """
            )

            facturas = cursor.fetchall()

            ingresos = sum(float(f[2] or 0) for f in facturas)

            return {
                "ingresos": ingresos,
                "ventas": len(facturas),
                "facturas": facturas
            }

        except Exception as error:
            print("Error al obtener los reportes:", error)
            return {
                "ingresos": 0,
                "ventas": 0,
                "facturas": []
            }

        finally:
            if cursor:
                cursor.close()
            conexion.close()

    def exportar_reporte_pdf(periodo_seleccionado):
        datos = obtener_datos_reportes(periodo_seleccionado)
        facturas = datos["facturas"]

        if not facturas:
            messagebox.showinfo(
                "Sin datos",
                "No hay facturas pagadas para exportar en este período."
            )
            return

        ruta = filedialog.asksaveasfilename(
            title="Guardar reporte PDF",
            defaultextension=".pdf",
            filetypes=[("Archivo PDF", "*.pdf")],
            initialfile=f"reporte_ventas_{periodo_seleccionado.lower().replace(' ', '_')}.pdf"
        )

        if not ruta:
            return

        try:
            documento = SimpleDocTemplate(
                ruta,
                pagesize=landscape(letter),
                rightMargin=30,
                leftMargin=30,
                topMargin=30,
                bottomMargin=30
            )

            estilos = getSampleStyleSheet()
            elementos = [
                Paragraph(
                    "Restaurante Vivaldi",
                    estilos["Title"]
                ),
                Paragraph(
                    f"Reporte de ventas cobradas: {periodo_seleccionado}",
                    estilos["Heading2"]
                ),
                Spacer(1, 15)
            ]

            tabla_datos = [
                ["N.º Factura", "Fecha", "Total", "Método de pago", "Mesero"]
            ]

            for factura in facturas:
                fecha = (
                    factura[1].strftime("%d/%m/%Y")
                    if hasattr(factura[1], "strftime")
                    else str(factura[1])
                )

                tabla_datos.append([
                    str(factura[0]),
                    fecha,
                    f"${float(factura[2] or 0):,.2f}",
                    str(factura[3] or "No especificado"),
                    str(factura[4] or "No asignado")
                ])

            tabla_datos.append([
                "",
                "TOTAL COBRADO",
                f"${datos['ingresos']:,.2f}",
                ""
            ])

            tabla = Table(
                tabla_datos,
                repeatRows=1,
                colWidths=[85, 110, 105, 150, 180]
            )

            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#223D2D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D9BC8C")),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9BC8C")),
                ("ALIGN", (2, 1), (2, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8)
            ]))

            elementos.append(tabla)
            documento.build(elementos)

            messagebox.showinfo(
                "Reporte generado",
                "El reporte PDF se guardó correctamente."
            )

        except Exception as error:
            messagebox.showerror(
                "Error al exportar PDF",
                f"No se pudo generar el reporte:\n{error}"
            )

    def exportar_reporte_excel(periodo_seleccionado):
        datos = obtener_datos_reportes(periodo_seleccionado)
        facturas = datos["facturas"]

        if not facturas:
            messagebox.showinfo(
                "Sin datos",
                "No hay facturas pagadas para exportar en este período."
            )
            return

        ruta = filedialog.asksaveasfilename(
            title="Guardar reporte Excel",
            defaultextension=".xlsx",
            filetypes=[("Archivo Excel", "*.xlsx")],
            initialfile=f"reporte_ventas_{periodo_seleccionado.lower().replace(' ', '_')}.xlsx"
        )

        if not ruta:
            return

        try:
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Ventas cobradas"

            hoja.append(["RESTAURANTE VIVALDI"])
            hoja.append([f"Reporte de ventas cobradas: {periodo_seleccionado}"])
            hoja.append([])
            hoja.append([
                "N.º Factura",
                "Fecha",
                "Total ($)",
                "Método de pago",
                "Mesero"
            ])

            for factura in facturas:
                hoja.append([
                    factura[0],
                    factura[1],
                    float(factura[2] or 0),
                    factura[3] or "No especificado",
                    factura[4] or "No asignado"
                ])

            fila_total = hoja.max_row + 1
            hoja.cell(row=fila_total, column=2, value="TOTAL COBRADO")
            hoja.cell(
                row=fila_total,
                column=3,
                value=datos["ingresos"]
            )

            for celda in hoja[4]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill(
                    "solid",
                    fgColor="223D2D"
                )
                celda.alignment = Alignment(horizontal="center")

            for fila in hoja.iter_rows(
                    min_row=5,
                    max_row=fila_total
            ):
                for celda in fila:
                    celda.alignment = Alignment(vertical="center")

            for fila in range(5, fila_total):
                hoja.cell(row=fila, column=3).number_format = '$#,##0.00'

            hoja.cell(row=fila_total, column=3).number_format = '$#,##0.00'

            for celda in hoja[fila_total]:
                celda.font = Font(bold=True)
                celda.fill = PatternFill(
                    "solid",
                    fgColor="D9BC8C"
                )

            hoja.column_dimensions["A"].width = 18
            hoja.column_dimensions["B"].width = 24
            hoja.column_dimensions["C"].width = 18
            hoja.column_dimensions["D"].width = 24
            hoja.column_dimensions["E"].width = 28

            hoja.freeze_panes = "A5"
            hoja.auto_filter.ref = f"A4:D{fila_total - 1}"

            libro.save(ruta)

            messagebox.showinfo(
                "Reporte generado",
                "El reporte Excel se guardó correctamente."
            )

        except Exception as error:
            messagebox.showerror(
                "Error al exportar Excel",
                f"No se pudo generar el reporte:\n{error}"
            )

    def obtener_resumen_reservaciones(periodo="Esta semana"):
        conexion = conectar()

        if not conexion:
            return {
                "reservaciones": 0,
                "clientes": 0,
                "estados": [
                    ("Confirmadas", 0, VERDE),
                    ("Pendientes", 0, BEIGE),
                    ("Canceladas", 0, VINO),
                    ("Finalizadas", 0, GRIS)
                ]
            }

        cursor = None

        try:
            cursor = conexion.cursor()

            filtros = {
                "Hoy": "fecha = CURDATE()",
                "Esta semana": (
                    "YEARWEEK(fecha, 1) = YEARWEEK(CURDATE(), 1)"
                ),
                "Este mes": (
                    "YEAR(fecha) = YEAR(CURDATE()) "
                    "AND MONTH(fecha) = MONTH(CURDATE())"
                ),
                "Este año": "YEAR(fecha) = YEAR(CURDATE())"
            }

            condicion = filtros.get(periodo, filtros["Esta semana"])

            cursor.execute(
                f"""
                SELECT
                    COUNT(*),
                    COALESCE(SUM(estado = 'Confirmada'), 0),
                    COALESCE(SUM(estado = 'Pendiente'), 0),
                    COALESCE(SUM(estado = 'Cancelada'), 0),
                    COALESCE(SUM(estado = 'Completada'), 0)
                FROM Reserva
                WHERE {condicion}
                """
            )

            fila = cursor.fetchone()

            reservaciones = int(fila[0] or 0)

            estados = [
                ("Confirmadas", int(fila[1] or 0), VERDE),
                ("Pendientes", int(fila[2] or 0), BEIGE),
                ("Canceladas", int(fila[3] or 0), VINO),
                ("Finalizadas", int(fila[4] or 0), GRIS)
            ]

            cursor.execute("SELECT COUNT(*) FROM Cliente")
            clientes = int(cursor.fetchone()[0] or 0)

            return {
                "reservaciones": reservaciones,
                "clientes": clientes,
                "estados": estados
            }

        except Exception as error:
            print("Error al obtener reservaciones y clientes:", error)
            raise

        finally:
            if cursor:
                cursor.close()
            conexion.close()

    def mostrar_reportes(periodo_inicial="Esta semana"):
        limpiar_vista()
        mostrar_encabezados(False)
        activar_boton("REPORTES")

        encabezado = tk.Frame(vista_actual, bg=CREMA)
        encabezado.pack(fill="x", pady=(8, 12))

        tk.Label(
            encabezado,
            text="Reportes y estadísticas",
            font=("Georgia", 22, "bold"),
            bg=CREMA,
            fg=VERDE_OSCURO
        ).pack(anchor="w")

        tk.Label(
            encabezado,
            text="Resumen del rendimiento de Restaurante Vivaldi",
            font=("Segoe UI", 9),
            bg=CREMA,
            fg=VERDE
        ).pack(anchor="w", pady=(5, 0))

        barra_filtros = tk.Frame(vista_actual, bg=CREMA)
        barra_filtros.pack(fill="x", pady=15)

        tk.Label(
            barra_filtros,
            text="Período del reporte:",
            font=("Arial", 10, "bold"),
            bg=CREMA,
            fg=VERDE_OSCURO
        ).pack(side="left", padx=(0, 10))

        periodo = tk.StringVar(value=periodo_inicial)

        datos = obtener_datos_reportes(periodo.get())
        ingresos_reales = datos["ingresos"]
        ventas_reales = datos["ventas"]

        resumen = obtener_resumen_reservaciones(periodo.get())
        reservaciones_reales = resumen["reservaciones"]
        clientes_reales = resumen["clientes"]

        selector_periodo = ttk.Combobox(
            barra_filtros,
            textvariable=periodo,
            values=[
                "Hoy",
                "Esta semana",
                "Este mes",
                "Este año"
            ],
            state="readonly",
            width=18
        )
        selector_periodo.pack(side="left", ipady=4)


        tk.Button(
            barra_filtros,
            text="Exportar PDF",
            command=lambda: exportar_reporte_pdf(periodo.get()),
            bg=VINO,
            fg=BLANCO,
            activebackground=VINO,
            activeforeground=BLANCO,
            font=("Arial", 9, "bold"),
            relief="flat",
            cursor="hand2",
            padx=12,
            pady=7
        ).pack(side="right", padx=(6, 0))

        tk.Button(
            barra_filtros,
            text="Exportar Excel",
            command=lambda: exportar_reporte_excel(periodo.get()),
            bg=VERDE,
            fg=BLANCO,
            activebackground=VERDE,
            activeforeground=BLANCO,
            font=("Arial", 9, "bold"),
            relief="flat",
            cursor="hand2",
            padx=12,
            pady=7
        ).pack(side="right", padx=(6, 0))


        contenedor = tk.Frame(vista_actual, bg=CREMA)
        contenedor.pack(
            fill="both",
            expand=True,
            padx=30,
            pady=(0, 25)
        )

        canvas = tk.Canvas(
            contenedor,
            bg=CREMA,
            highlightthickness=0
        )

        scrollbar = ttk.Scrollbar(
            contenedor,
            orient="vertical",
            command=canvas.yview
        )

        canvas.configure(yscrollcommand=scrollbar.set)

        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        contenido = tk.Frame(canvas, bg=CREMA)

        ventana_contenido = canvas.create_window(
            (0, 0),
            window=contenido,
            anchor="nw"
        )

        def actualizar_scroll(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def ajustar_ancho(event):
            canvas.itemconfigure(
                ventana_contenido,
                width=event.width
            )

        contenido.bind("<Configure>", actualizar_scroll)
        canvas.bind("<Configure>", ajustar_ancho)

        def desplazar_reportes(event):
            canvas.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units"
            )

        canvas.bind("<MouseWheel>", desplazar_reportes)

        indicadores = tk.Frame(contenido, bg=CREMA)
        indicadores.pack(fill="x", pady=(8, 20))

        datos_indicadores = [
            ("INGRESOS", f"${ingresos_reales:,.2f}", VERDE),
            ("VENTAS COBRADAS", str(ventas_reales), VERDE_OSCURO),
            ("RESERVACIONES", str(reservaciones_reales), VINO),
            ("CLIENTES", str(clientes_reales), VERDE_OSCURO)
        ]

        for columna, (titulo, valor, color) in enumerate(datos_indicadores):
            tarjeta = tk.Frame(
                indicadores,
                bg=BLANCO,
                highlightbackground=BEIGE,
                highlightthickness=1
            )

            tarjeta.grid(
                row=0,
                column=columna,
                sticky="nsew",
                padx=6,
                pady=5
            )

            tk.Frame(
                tarjeta,
                bg=color,
                height=5
            ).pack(fill="x")

            tk.Label(
                tarjeta,
                text=titulo,
                font=("Arial", 9, "bold"),
                bg=BLANCO,
                fg=GRIS
            ).pack(anchor="w", padx=12, pady=(14, 5))

            tk.Label(
                tarjeta,
                text=valor,
                font=("Georgia", 19, "bold"),
                bg=BLANCO,
                fg=VERDE_OSCURO
            ).pack(anchor="w", pady=(8, 12))

            indicadores.grid_columnconfigure(
                columna,
                weight=1,
                uniform="indicadores"
            )

        fila_graficos = tk.Frame(contenido, bg=CREMA)
        fila_graficos.pack(fill="x", pady=(8, 12))

        tarjeta_ingresos = tk.Frame(
            fila_graficos,
            bg=BLANCO,
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        tarjeta_ingresos.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(0, 8)
        )

        tk.Label(
            tarjeta_ingresos,
            text="Ingresos semanales",
            font=("Georgia", 15, "bold"),
            bg=BLANCO,
            fg=VERDE_OSCURO
        ).pack(anchor="w", pady=(8, 12))

        tk.Label(
            tarjeta_ingresos,
            text="Ingresos registrados por día",
            font=("Arial", 9),
            bg=BLANCO,
            fg=GRIS
        ).pack(anchor="w", padx=18)

        grafica_ingresos = tk.Canvas(
            tarjeta_ingresos,
            bg=BLANCO,
            height=250,
            highlightthickness=0
        )
        grafica_ingresos.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

        def obtener_ingresos_grafica(periodo_seleccionado):
            conexion = conectar()

            if not conexion:
                return []

            cursor = None

            try:
                cursor = conexion.cursor()

                if periodo_seleccionado == "Hoy":
                    agrupacion = "HOUR(fecha)"
                    etiqueta = "DATE_FORMAT(fecha, '%H:00')"
                    condicion = "DATE(fecha) = CURDATE()"
                elif periodo_seleccionado == "Esta semana":
                    agrupacion = "DATE(fecha)"
                    etiqueta = "DATE_FORMAT(fecha, '%a %d')"
                    condicion = (
                        "YEARWEEK(fecha, 1) = YEARWEEK(CURDATE(), 1)"
                    )
                elif periodo_seleccionado == "Este mes":
                    agrupacion = "DATE(fecha)"
                    etiqueta = "DATE_FORMAT(fecha, '%d/%m')"
                    condicion = (
                        "YEAR(fecha) = YEAR(CURDATE()) "
                        "AND MONTH(fecha) = MONTH(CURDATE())"
                    )
                else:
                    agrupacion = "MONTH(fecha)"
                    etiqueta = "DATE_FORMAT(fecha, '%b')"
                    condicion = "YEAR(fecha) = YEAR(CURDATE())"

                cursor.execute(
                    f"""
                    SELECT {etiqueta}, SUM(total)
                    FROM Factura
                    WHERE estado_pago = 'Pagado'
                      AND {condicion}
                    GROUP BY {agrupacion}, {etiqueta}
                    ORDER BY {agrupacion}
                    """.replace(
                        "SELECT " + etiqueta,
                        "SELECT " + etiqueta,
                        1
                    )
                )

                return [
                    (str(fila[0]), float(fila[1] or 0))
                    for fila in cursor.fetchall()
                ]

            except Exception as error:
                print("Error al obtener ingresos para el gráfico:", error)
                return []

            finally:
                if cursor:
                    cursor.close()
                conexion.close()

        ingresos = obtener_ingresos_grafica(periodo.get())

        def dibujar_ingresos(event=None):
            grafica_ingresos.delete("all")

            ancho = grafica_ingresos.winfo_width()
            alto = grafica_ingresos.winfo_height()

            if ancho < 100 or alto < 100:
                return

            margen_izquierdo = 35
            margen_inferior = 35
            margen_superior = 20

            altura_util = alto - margen_inferior - margen_superior
            ancho_util = ancho - margen_izquierdo - 15
            maximo = max(
                (cantidad for _, cantidad in ingresos),
                default=0
            )
            maximo = maximo * 1.15 if maximo > 0 else 1

            for indice in range(6):
                cantidad = maximo * indice / 5
                y = (
                    alto - margen_inferior
                    - (cantidad / maximo) * altura_util
                )

                grafica_ingresos.create_line(
                    margen_izquierdo,
                    y,
                    ancho - 10,
                    y,
                    fill="#E8E0D2"
                )

                grafica_ingresos.create_text(
                    margen_izquierdo - 5,
                    y,
                    text=str(cantidad),
                    anchor="e",
                    font=("Arial", 8),
                    fill=GRIS
                )

                if not ingresos:
                    grafica_ingresos.create_text(
                        margen_izquierdo // 2,
                        margen_superior // 2,
                        text="No hay ventas registradas en este período",
                        font=("Arial", 11),
                        fill=GRIS
                    )
                    return

            espacio = ancho_util / len(ingresos)
            ancho_barra = min(28, espacio * 0.55)

            for indice, (dia, cantidad) in enumerate(ingresos):
                centro_x = (
                    margen_izquierdo
                    + espacio * indice
                    + espacio / 2
                )

                altura_barra = (cantidad / maximo) * altura_util
                y_superior = alto - margen_inferior - altura_barra

                grafica_ingresos.create_rectangle(
                    centro_x - ancho_barra / 2,
                    y_superior,
                    centro_x + ancho_barra / 2,
                    alto - margen_inferior,
                    fill=VERDE,
                    outline=""
                )

                grafica_ingresos.create_text(
                    centro_x,
                    y_superior - 10,
                    text=f"${cantidad}",
                    font=("Arial", 8, "bold"),
                    fill=VERDE_OSCURO
                )

                grafica_ingresos.create_text(
                    centro_x,
                    alto - margen_inferior + 16,
                    text=dia,
                    font=("Arial", 8),
                    fill=GRIS
                )

        grafica_ingresos.bind(
            "<Configure>",
            dibujar_ingresos
        )

        # Gráfico de reservaciones
        tarjeta_reservaciones = tk.Frame(
            fila_graficos,
            bg=BLANCO,
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        tarjeta_reservaciones.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(8, 0)
        )

        tk.Label(
            tarjeta_reservaciones,
            text="Estado de reservaciones",
            font=("Georgia", 15, "bold"),
            bg=BLANCO,
            fg=VERDE_OSCURO
        ).pack(anchor="w", padx=18, pady=(18, 3))

        tk.Label(
            tarjeta_reservaciones,
            text="Distribución de las reservaciones",
            font=("Arial", 9),
            bg=BLANCO,
            fg=GRIS
        ).pack(anchor="w", padx=18)

        grafica_reservaciones = tk.Canvas(
            tarjeta_reservaciones,
            bg=BLANCO,
            height=250,
            highlightthickness=0
        )
        grafica_reservaciones.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

        estados = resumen["estados"]

        def dibujar_reservaciones(event=None):
            grafica_reservaciones.delete("all")

            ancho = grafica_reservaciones.winfo_width()

            if ancho < 100:
                return

            total = sum(cantidad for _, cantidad, _ in estados)

            y = 30
            ancho_barra = ancho - 45

            for nombre, cantidad, color in estados:
                grafica_reservaciones.create_text(
                    12,
                    y,
                    text=nombre,
                    anchor="w",
                    font=("Arial", 9, "bold"),
                    fill=VERDE_OSCURO
                )

                grafica_reservaciones.create_text(
                    ancho - 12,
                    y,
                    text=f"{cantidad} ({cantidad / total * 100:.0f}%)",
                    anchor="e",
                    font=("Arial", 9),
                    fill=GRIS
                )

                y += 12

                grafica_reservaciones.create_rectangle(
                    12,
                    y,
                    12 + ancho_barra,
                    y + 13,
                    fill="#EEE8DD",
                    outline=""
                )

                grafica_reservaciones.create_rectangle(
                    12,
                    y,
                    12 + ancho_barra * cantidad / total,
                    y + 13,
                    fill=color,
                    outline=""
                )

                y += 42

        grafica_reservaciones.bind(
            "<Configure>",
            dibujar_reservaciones
        )


        # Resumen del período
        tarjeta_resumen = tk.Frame(
            contenido,
            bg=BLANCO,
            highlightbackground=BEIGE,
            highlightthickness=1
        )
        tarjeta_resumen.pack(fill="x", pady=(0, 15))

        tk.Label(
            tarjeta_resumen,
            text="Resumen del período",
            font=("Georgia", 15, "bold"),
            bg=BLANCO,
            fg=VERDE_OSCURO
        ).pack(anchor="w", padx=20, pady=(18, 10))

        texto_resumen = (
            f"Período consultado: {periodo.get()}\n"
            f"Ingresos cobrados: ${ingresos_reales:,.2f}\n"
            f"Cantidad de facturas pagadas: {ventas_reales}\n"
            f"Reservaciones registradas: {reservaciones_reales}\n"
            f"Clientes registrados en el sistema: {clientes_reales}"
        )

        tk.Label(
            tarjeta_resumen,
            text=texto_resumen,
            font=("Arial", 10),
            bg=BLANCO,
            fg=GRIS,
            justify="left",
            wraplength=850
        ).pack(anchor="w", padx=20, pady=(0, 18))

        def cambiar_periodo(event=None):
            periodo_seleccionado = periodo.get()
            mostrar_reportes(periodo_seleccionado)

        selector_periodo.bind(
            "<<ComboboxSelected>>",
            cambiar_periodo
        )

    def cerrar_sesion():
        respuesta = messagebox.askyesno(
            "Cerrar sesión",
            "¿Deseas cerrar sesión?"
        )

        if not respuesta:
            return

        try:
            ruta_login = os.path.join(
                os.path.dirname(os.path.abspath(__file__)),
                "inicio_sistema.py"
            )

            subprocess.Popen([
                sys.executable,
                ruta_login
            ])

            ventana.destroy()

        except Exception as error:
            messagebox.showerror(
                "Error al cerrar sesión",
                f"No se pudo abrir el inicio de sesión:\n{error}"
            )

    crear_boton_menu(
        "INICIO",
        mostrar_inicio
    )

    crear_boton_menu(
        "MESAS",
        mostrar_mesas
    )

    crear_boton_menu(
        "MENÚ",
        mostrar_menu
    )

    crear_boton_menu(
        "INVENTARIO",
        mostrar_inventario
    )

    crear_boton_menu(
        "PROVEEDORES",
        mostrar_proveedores
    )

    crear_boton_menu(
        "REPORTES",
        mostrar_reportes
    )

    tk.Frame(
        menu_navegacion,
        bg=BEIGE,
        height=1
    ).pack(fill="x", pady=(18, 10))

    tk.Button(
        menu_navegacion,
        text="CERRAR SESIÓN",
        bg=VERDE_OSCURO,
        fg=CREMA_CLARO,
        activebackground=CREMA_CLARO,
        activeforeground=BLANCO,
        bd=0,
        relief="flat",
        anchor="w",
        padx=10,
        pady=13,
        font=("Arial", 10, "bold"),
        cursor="hand2",
        command=cerrar_sesion
    ).pack(fill="x", pady=4)

    mostrar_inicio()
    ventana.mainloop()

if __name__ == "__main__":

    interfaz_administrador()